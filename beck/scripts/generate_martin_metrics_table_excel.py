from __future__ import annotations

import sys
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import Reference, ScatterChart, Series
from openpyxl.styles import Alignment, Font


CURRENT_DIR = Path(__file__).resolve().parent
if str(CURRENT_DIR) not in sys.path:
    sys.path.insert(0, str(CURRENT_DIR))

from generate_martin_arch_quality_doc import (  # noqa: E402
    TARGET_LAYERS,
    collect_modules,
    compute_metrics,
)


OUT_PATH = r"e:\sem8\ful\Diplom_NK\beck\docs\martin_metrics_table.xlsx"


def _autosize(ws) -> None:
    for col in ws.columns:
        max_len = 0
        letter = col[0].column_letter
        for cell in col:
            v = "" if cell.value is None else str(cell.value)
            max_len = max(max_len, len(v))
        ws.column_dimensions[letter].width = min(28, max(10, max_len + 2))


def main() -> None:
    modules = collect_modules()
    metrics = compute_metrics(modules)

    wb = Workbook()
    ws = wb.active
    ws.title = "A-I-D Table"

    ws["A1"] = "Таблица метрик Роберта Мартина"
    ws["A2"] = f"Дата расчета: {datetime.now().strftime('%d.%m.%Y %H:%M')}"
    ws["A1"].font = Font(size=14, bold=True)
    ws["A2"].font = Font(italic=True)

    headers = ["Слой", "A", "I", "D", "Nc", "Na", "Ca", "Ce"]
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=4, column=i, value=h)
        c.font = Font(bold=True)
        c.alignment = Alignment(horizontal="center")

    row = 5
    for layer in TARGET_LAYERS:
        m = metrics[layer]
        ws.cell(row=row, column=1, value=layer)
        ws.cell(row=row, column=2, value=round(m["A"], 4))
        ws.cell(row=row, column=3, value=round(m["I"], 4))
        ws.cell(row=row, column=4, value=round(m["D"], 4))
        ws.cell(row=row, column=5, value=m["Nc"])
        ws.cell(row=row, column=6, value=m["Na"])
        ws.cell(row=row, column=7, value=m["Ca"])
        ws.cell(row=row, column=8, value=m["Ce"])
        row += 1

    ws["A10"] = "Формулы:"
    ws["A10"].font = Font(bold=True)
    ws["A11"] = "A = Na / Nc"
    ws["A12"] = "I = Ce / (Ca + Ce)"
    ws["A13"] = "D = |A + I - 1|"

    # Данные для диагонали главной последовательности: (0,1) -> (1,0)
    ws["J4"] = "X_diag"
    ws["K4"] = "Y_diag"
    ws["J5"] = 0.0
    ws["K5"] = 1.0
    ws["J6"] = 1.0
    ws["K6"] = 0.0

    # Scatter-график A/I: точки слоев + диагональ Main Sequence.
    chart = ScatterChart()
    chart.title = "Main Sequence: точки A/I"
    chart.x_axis.title = "A (Abstractness)"
    chart.y_axis.title = "I (Instability)"
    chart.width = 14
    chart.height = 7

    # Оси строго от 0 до 1.
    chart.x_axis.scaling.min = 0
    chart.x_axis.scaling.max = 1
    chart.y_axis.scaling.min = 0
    chart.y_axis.scaling.max = 1

    # Точки (x=A, y=I) по слоям.
    x_points = Reference(ws, min_col=2, min_row=5, max_row=4 + len(TARGET_LAYERS))
    y_points = Reference(ws, min_col=3, min_row=5, max_row=4 + len(TARGET_LAYERS))
    points = Series(y_points, x_points, title="Слои (A,I)")
    points.marker.symbol = "circle"
    points.marker.size = 9
    points.graphicalProperties.line.noFill = True
    chart.series.append(points)

    # Диагональ от (0,1) до (1,0).
    x_diag = Reference(ws, min_col=10, min_row=5, max_row=6)
    y_diag = Reference(ws, min_col=11, min_row=5, max_row=6)
    diag = Series(y_diag, x_diag, title="Линия I = 1 - A")
    diag.marker.symbol = "none"
    chart.series.append(diag)

    ws.add_chart(chart, "J8")

    _autosize(ws)
    wb.save(OUT_PATH)
    print(OUT_PATH)


if __name__ == "__main__":
    main()
