from datetime import datetime

from docx import Document
from docx.shared import Pt


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_numbers(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Number")


def main() -> None:
    out_path = r"e:\sem8\ful\Diplom_NK\beck\docs\architecture_beck.docx"

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    doc.add_heading("Архитектурный обзор проекта beck", 0)
    doc.add_paragraph(f"Дата формирования: {datetime.now().strftime('%d.%m.%Y %H:%M')}.")
    doc.add_paragraph(
        "Документ описывает текущую архитектуру проекта, ключевые потоки данных, "
        "используемые паттерны проектирования, технологические решения, риски и "
        "рекомендованный план эволюции архитектуры."
    )

    doc.add_heading("1. Назначение системы и контекст", level=1)
    doc.add_paragraph(
        "Система реализует генерацию и автоматическое дообогащение технологической карты "
        "радиографического контроля сварных соединений. Пользователь выбирает параметры "
        "в карточке (тип сварного соединения, материал, ИИИ, схема просвечивания и др.), "
        "после чего сервер последовательно применяет набор бизнес-правил (changers), "
        "получает справочные данные из PostgreSQL и рассчитывает зависимые поля."
    )
    doc.add_paragraph(
        "Архитектурный стиль: layered monolith с явным pipeline-движком правил и "
        "репозиторием для доступа к БД. Система ориентирована на расширение через добавление "
        "новых changer-ов без изменений базового orchestration-кода."
    )

    doc.add_heading("2. Структура проекта по слоям", level=1)
    add_bullets(
        doc,
        [
            "main.py — composition root: сборка зависимостей, pipeline и запуск web-адаптера.",
            "controllers/ — транспортный слой (FastAPI endpoint, маршрутизация control_type, обработка ошибок).",
            "services/ — прикладной слой: модель карточки TechCardData, сервис use-case, Pipeline, changers.",
            "services/Changers/ — набор бизнес-правил и вычислителей.",
            "repositories/ — инфраструктурный слой: SQL к PostgreSQL и сохранение/загрузка снапшотов карт.",
            "repositories/Interfaces/ — портовый слой (контракты для data access).",
        ],
    )

    doc.add_heading("3. Ключевые компоненты", level=1)
    doc.add_heading("3.1 Entry point и DI", level=2)
    doc.add_paragraph(
        "В main.py формируется единый граф зависимостей: PostgresDataBase -> TechCardService -> "
        "ControllerWeb -> adapterWeb. В том же месте фиксируется порядок changer-ов в pipeline. "
        "Это фактический центр архитектурной конфигурации."
    )

    doc.add_heading("3.2 Transport layer", level=2)
    doc.add_paragraph(
        "adapterWeb.py поднимает HTTP endpoint /techcard/{control_type}. "
        "Значение control_type определяет сценарий: template, updateTechCard, saveTechCard, "
        "listSavedTechCards, getSavedTechCard. ControllerWeb остаётся тонким фасадом и делегирует "
        "операции в TechCardService."
    )

    doc.add_heading("3.3 Application layer", level=2)
    doc.add_paragraph(
        "TechCardService управляет жизненным циклом карточки: построение шаблона, запуск pipeline, "
        "подготовка данных для сохранения и чтения снапшотов. Класс TechCardData является mutable-model "
        "со специализированными операциями по вставке/обновлению/удалению параметров в блоках."
    )

    doc.add_heading("3.4 Rule engine (Pipeline + Changers)", level=2)
    doc.add_paragraph(
        "PipeLine выполняет changers последовательно. Каждый changer реализует IDataChanger и "
        "инкапсулирует конкретную бизнес-задачу: загрузка справочника, вычисление по формуле, "
        "синхронизация зависимых полей, валидация доступности вариантов."
    )

    doc.add_heading("3.5 Data access", level=2)
    doc.add_paragraph(
        "PostgresDataBase объединяет реализацию нескольких интерфейсов репозитория: "
        "справочники, параметры сварного соединения, аппараты, плёнки, привязка материала "
        "к стандарту, таблица напряжения на трубке (voltage_tube), операции сохранения и "
        "чтения карточек."
    )

    doc.add_heading("4. Runtime-потоки", level=1)
    doc.add_heading("4.1 Получение шаблона", level=2)
    add_numbers(
        doc,
        [
            "POST /techcard/template -> controller.get_template().",
            "TechCardService формирует базовую структуру карты и блок операций.",
            "Pipeline последовательно донастраивает поля, options и значения.",
            "Результат возвращается клиенту как сериализованная структура карточки.",
        ],
    )

    doc.add_heading("4.2 Обновление карты", level=2)
    add_numbers(
        doc,
        [
            "POST /techcard/updateTechCard принимает текущую карту от клиента.",
            "ControllerWeb десериализует payload в TechCardData.",
            "Pipeline повторно прогоняет все правила и синхронизирует зависимые параметры.",
            "Клиент получает консистентную карточку после применения бизнес-правил.",
        ],
    )

    doc.add_heading("5. Активные бизнес-правила (changers)", level=1)
    add_bullets(
        doc,
        [
            "ControlMethodsFromDb, RegulatoryDocumentsFromDb — загрузка справочников методики и нормативов.",
            "TypeOfWeldedJointFromDb, WeldedJointDiagramFromDb — выбор типа шва и соответствующей схемы.",
            "ParamsByWeldedJointFromDb — синхронизация параметров по типу шва с флагом источника.",
            "CircButtWeldBeadFromThicknessTable — расчет e/g (или A,h1,h2) по таблице толщины.",
            "RengenApparatusFromDb и RengenApparatusForPanoramicScheme — выбор ИИИ с ограничениями и вычислением сопутствующих полей.",
            "xray_max_kv_by_material_table + repository voltage_tube — расчет максимально допустимого кВ из БД.",
            "ControlSensitivityChanger, RadiographicFilmFromDb, IntensifyingScreenByVoltage, ProtectiveScreenByVoltage — расчет чувствительности, пленки и экранов.",
            "RengenDistanceForDoubleWallScheme / EllipseFrontal / RectilinearDrDistance — формулы расстояний для разных схем.",
            "ControlledZoneWidthStub — вычисление ширины контролируемой зоны по толщине.",
        ],
    )

    doc.add_heading("6. Применяемые паттерны проектирования", level=1)
    patterns = [
        (
            "Pipeline / Chain of Responsibility",
            "PipeLine запускает последовательность IDataChanger. Каждый шаг применяет узкое правило и изменяет одну модель данных.",
        ),
        (
            "Strategy",
            "Каждый changer является независимой стратегией обработки с единым контрактом changeData().",
        ),
        (
            "Repository",
            "SQL изолирован в PostgresDataBase и вызывается через интерфейсы i_*.py.",
        ),
        (
            "Dependency Injection (manual)",
            "Зависимости передаются через конструкторы; составление графа в main.py.",
        ),
        (
            "Adapter",
            "adapterWeb преобразует HTTP payload и control_type в вызовы контроллера/сервиса.",
        ),
        (
            "Facade",
            "TechCardService предоставляет единый API use-case поверх pipeline и persistence.",
        ),
    ]
    for name, desc in patterns:
        p = doc.add_paragraph()
        p.add_run(name + ": ").bold = True
        p.add_run(desc)

    doc.add_heading("7. Технические сильные стороны", level=1)
    add_bullets(
        doc,
        [
            "Четкое разделение ответственности между слоями transport/application/data.",
            "Высокая расширяемость бизнес-логики через добавление новых changer-ов.",
            "Накопление доменных формул в отдельных модулях, а не в контроллерах.",
            "Явные интерфейсы репозитория облегчают замену источника данных и тестирование.",
            "Идемпотентные обновления большинства параметров снижают вероятность дублей в карте.",
        ],
    )

    doc.add_heading("8. Ограничения и архитектурные риски", level=1)
    add_bullets(
        doc,
        [
            "PostgresDataBase содержит слишком много ответственности (God Object в data layer).",
            "Поведение сильно order-dependent: изменение порядка changer-ов может вызывать регрессии.",
            "Широкое использование строковых имен полей/блоков повышает хрупкость.",
            "Есть смешение современных и legacy-подходов в части changer-ов и утилит.",
            "DSN с паролем в main.py — риск безопасности и переносимости конфигурации.",
            "Ограниченная наблюдаемость: print-логирование вместо структурированного логирования/метрик.",
        ],
    )

    doc.add_heading("9. Архитектура работы с напряжением трубки (актуальный кейс)", level=1)
    doc.add_paragraph(
        "В проекте используется таблица public.voltage_tube, связанная с type_metall. "
        "Предел кВ вычисляется по материалу и радиационной толщине. При отсутствии данных "
        "предел считается равным 0 кВ (fail-safe fallback). Это значение записывается в "
        "параметр «Напряжение на рентгеновской трубке, не более, кВ» и участвует в фильтрации "
        "ИИИ. Логика фильтрации учитывает, что аппарат допустим только если его максимальное "
        "напряжение не выше допустимого предела."
    )

    doc.add_heading("10. Рекомендации по эволюции архитектуры", level=1)
    doc.add_heading("10.1 Краткосрочно (1-3 недели)", level=2)
    add_bullets(
        doc,
        [
            "Вынести DB_DSN и секреты в переменные окружения.",
            "Добавить интеграционные тесты на критические пользовательские сценарии pipeline.",
            "Унифицировать константы имен блоков и параметров в одном модуле.",
            "Добавить unit-тесты на репозиторные lookup-методы (в т.ч. voltage_tube).",
        ],
    )

    doc.add_heading("10.2 Среднесрочно (1-2 месяца)", level=2)
    add_bullets(
        doc,
        [
            "Разделить PostgresDataBase на несколько специализированных repository-классов.",
            "Ввести typed DTO/Value Object для ключевых полей вместо произвольных dict.",
            "Формализовать зависимости между changer-ами (декларация prerequisite-полей).",
            "Стабилизировать API-контракт ответов (единый формат сериализации).",
        ],
    )

    doc.add_heading("10.3 Долгосрочно (2-6 месяцев)", level=2)
    add_bullets(
        doc,
        [
            "Выделить модуль rule-engine с версионированием правил и трассировкой источника значения.",
            "Добавить observability: структурированные логи, метрики времени шагов pipeline, аудит изменений полей.",
            "Постепенно перейти от stringly-typed модели к схемам с валидацией и обратной совместимостью.",
        ],
    )

    doc.add_heading("11. Итог", level=1)
    doc.add_paragraph(
        "Текущая архитектура подходит для инженерного продукта с большим числом доменных "
        "правил и частыми изменениями требований. Главная сильная сторона — расширяемость "
        "через changers и прозрачный orchestration в pipeline. Ключевые точки роста — "
        "декомпозиция data layer, типобезопасность модели, тестовое покрытие и стандартизация "
        "операционных практик."
    )

    doc.save(out_path)
    print(out_path)


if __name__ == "__main__":
    main()
