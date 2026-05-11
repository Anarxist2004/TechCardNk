from __future__ import annotations

from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.shared import Pt


ROOT = Path(r"e:\sem8\ful\Diplom_NK\beck")
OUTPUT = ROOT / "docs" / "rectilinear_dr_distance_detailed.docx"


def _setup(doc: Document) -> None:
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(14)
    pf = style.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_after = Pt(6)


def _p(doc: Document, text: str, bullet: bool = False) -> None:
    doc.add_paragraph(text, style="List Bullet" if bullet else None)


def _h(doc: Document, text: str, level: int) -> None:
    doc.add_heading(text, level=level)


def main() -> None:
    doc = Document()
    _setup(doc)

    doc.add_heading(
        "Подробное описание changer RectilinearDrDistance",
        0,
    )
    _p(
        doc,
        f"Документ подготовлен по модулю "
        f"`services/Changers/ch_RectilinearDrDistance.py`. "
        f"Дата: {datetime.now().strftime('%d.%m.%Y %H:%M')}.",
    )
    _p(
        doc,
        "Назначение changer-а: для прямолинейного сварного соединения в цифровой "
        "радиографии рассчитывать минимально допустимое расстояние f от источника "
        "излучения до поверхности объекта контроля и предельную длину участка "
        "экспонирования Lуч.",
    )

    _h(doc, "1. Обозначения и входные параметры", 1)
    for line in [
        "t — радиационная толщина, мм (поле «Радиационная толщина, мм»).",
        "Φ — размер фокусного пятна ИИИ, мм.",
        "K — чувствительность контроля, мм.",
        "Q — уровень качества изображения (A/B/C).",
        "S_img — коэффициент класса изображения: A->1.2, B->1.1, C->1.5.",
        "n — коэффициент, зависящий от Q и t (табличная piecewise-функция).",
        "b — расстояние от ППД до поверхности металла, мм (если поле найдено; иначе 0).",
        "t_sum — эффективная толщина в формуле: t + b.",
        "f_sto — значение f по основной формуле СТО Газпром.",
        "f_four — нижний страховочный предел 4*(t+b).",
        "f_min — итоговая минимальная граница расстояния.",
        "Lуч_max — предельная длина участка экспонирования.",
    ]:
        _p(doc, line, bullet=True)

    _h(doc, "2. Условия активации", 1)
    _p(
        doc,
        "Расчет активируется только для схем, распознаваемых как "
        "«Прямолинейное сварное соединение» (по фрагментам названия схемы).",
    )
    for line in [
        "Требуется блок «ИСХОДНЫЕ ДАННЫЕ».",
        "Должны быть валидны t, Φ, K и Q.",
        "Ограничения: t > 0, K > 0.",
        "Если условия не выполнены — placeholder/hint расстояния очищается (None).",
    ]:
        _p(doc, line, bullet=True)

    _h(doc, "3. Формулы и инженерная модель", 1)
    _h(doc, "3.1 Основная формула", 2)
    _p(doc, "Согласно реализованной логике (СТО Газпром, разд. 9.5.x):")
    _p(doc, "f = c * S_img * (t + b), где c = n * Φ / K", bullet=True)
    _p(doc, "Эквивалентная форма:")
    _p(doc, "f_sto = n * Φ * S_img * (t + b) / K", bullet=True)

    _h(doc, "3.2 Страховочное ограничение", 2)
    _p(doc, "Дополнительно применяется нижняя граница:")
    _p(doc, "f_four = 4 * (t + b)", bullet=True)
    _p(doc, "Итоговое минимальное расстояние:")
    _p(doc, "f_min = max(f_sto, f_four)", bullet=True)

    _h(doc, "3.3 Ограничение по длине участка экспонирования", 2)
    _p(doc, "В коде дополнительно рассчитывается:")
    _p(doc, "Lуч_max = 0.8 * f_min", bullet=True)
    _p(
        doc,
        "Подсказка в карту выводится в формате "
        "«f≥...; Lуч≤... (п. 9.5.21)».",
    )

    _h(doc, "4. Табличная логика коэффициента n", 1)
    _p(
        doc,
        "Функция _n_phi_over_k_divisor задает n в зависимости от Q и t:",
    )
    for line in [
        "Q=A: n=2 при t<50; n=3 при 50<=t<=100; n=4 при t>100.",
        "Q=B: n=2 при t<=100; n=3 при t>100.",
        "Q=C: n=2.",
    ]:
        _p(doc, line, bullet=True)

    _h(doc, "5. Пошаговый алгоритм работы changer-а", 1)
    for step in [
        "Проверка, что схема относится к прямолинейному соединению.",
        "Чтение t, Φ, K, Q из параметров карты.",
        "Поиск параметра b по текстовым фрагментам («детектор», «поверхност»); если не найден — b=0.",
        "Проверка валидности входных данных (t>0, K>0 и т.д.).",
        "Определение n по функции _n_phi_over_k_divisor(Q, t).",
        "Определение S_img по классу качества (A/B/C).",
        "Расчет t_sum = t + b.",
        "Расчет f_sto = n*Φ*S_img*t_sum/K.",
        "Расчет f_four = 4*t_sum.",
        "Расчет f_min = max(f_sto, f_four).",
        "Расчет Lуч_max = 0.8*f_min.",
        "Запись итоговой текстовой подсказки в placeholder/hint.",
    ]:
        _p(doc, step, bullet=True)

    _h(doc, "6. Поведение при неполных данных", 1)
    _p(
        doc,
        "Модуль использует безопасную стратегию: при невозможности надежного расчета "
        "не оставляет потенциально некорректную рекомендацию.",
    )
    for line in [
        "Если схема не прямолинейная — расчет не выполняется.",
        "Если отсутствуют/невалидны t, Φ, K, Q — подсказка очищается.",
        "Если t_sum<=0 или n/S_img не определены — подсказка очищается.",
        "Параметр b опционален: при отсутствии принимается b=0, что позволяет не блокировать расчет.",
    ]:
        _p(doc, line, bullet=True)

    _h(doc, "7. Что записывается в карту", 1)
    for line in [
        "Изменяется поле расстояния от ИИИ до поверхности контролируемого соединения.",
        "Обновляются оба атрибута: placeholder и hint.",
        "Формат: «f≥X; Lуч≤Y (п. 9.5.21)».",
    ]:
        _p(doc, line, bullet=True)

    _h(doc, "8. Вывод для ВКР", 1)
    _p(
        doc,
        "RectilinearDrDistance реализует формализованный инженерный расчет "
        "для прямолинейного шва в цифровой радиографии с учетом качества изображения, "
        "чувствительности, геометрических условий и нормативного страховочного ограничения. "
        "Это делает модуль ключевым для автоматизированной выдачи корректных "
        "технологических рекомендаций в рамках ОТК.",
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(str(OUTPUT))


if __name__ == "__main__":
    main()
