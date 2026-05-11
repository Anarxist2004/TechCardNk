from datetime import datetime

from docx import Document
from docx.shared import Pt


OUT_PATH = r"e:\sem8\ful\Diplom_NK\beck\docs\db_description_for_thesis.docx"


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_numbers(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Number")


def main() -> None:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    doc.add_heading("Описание базы данных системы радиографического контроля", 0)
    doc.add_paragraph(f"Дата формирования: {datetime.now().strftime('%d.%m.%Y %H:%M')}")
    doc.add_paragraph(
        "Документ подготовлен в формате пояснительной записки для дипломной работы. "
        "В нем подробно рассмотрены структура и логика предметно-ориентированной базы "
        "данных, используемой для формирования технологических карт радиографического "
        "контроля сварных соединений."
    )

    doc.add_heading("1. Назначение базы данных", level=1)
    doc.add_paragraph(
        "База данных предназначена для поддержки автоматизированного подбора параметров "
        "радиографического контроля. Хранилище объединяет нормативные справочники, "
        "технологические таблицы выбора, связи между типами сварных соединений и схемами "
        "контроля, а также архив сформированных техкарт в формате JSONB."
    )
    doc.add_paragraph(
        "Ключевая идея модели: пользователь работает с единой карточкой контроля, а сервер "
        "на лету вычисляет и подставляет необходимые параметры, используя набор таблиц БД "
        "как источник экспертных правил."
    )

    doc.add_heading("2. Логическая декомпозиция предметной области", level=1)
    add_bullets(
        doc,
        [
            "Подсистема нормативов и методик: control_methods, regulatory_documents.",
            "Подсистема сварного соединения: type_of_welded_joint, params_by_type_welding_joint, cheme_control, weld_type_to_scheme, width_of_controlled_area.",
            "Подсистема материалов и источников: type_metall, voltage_tube, class_film_radiation_source, rengen_apparatus, radiographic_film.",
            "Подсистема операций: list_operation, operation_param.",
            "Подсистема сохранения пользовательских результатов: tech_cards.",
            "Дополнительные технологические справочники: groove_standard, groove_standard_diam, control_sensitivity.",
        ],
    )
    doc.add_paragraph(
        "Такая декомпозиция отражает реальный жизненный цикл данных: от фиксированных "
        "нормативов к расчетным выборам и далее к сохранению готового состояния карты."
    )

    doc.add_heading("3. Подробный разбор таблиц", level=1)

    doc.add_heading("3.1 Нормативный контур", level=2)
    doc.add_paragraph("Таблица control_methods")
    add_bullets(
        doc,
        [
            "Содержит перечень методов контроля.",
            "PK: id.",
            "Является родительской сущностью для regulatory_documents.",
        ],
    )
    doc.add_paragraph("Таблица regulatory_documents")
    add_bullets(
        doc,
        [
            "Содержит перечень нормативных документов.",
            "PK: id.",
            "FK: control_methods_id -> control_methods.id.",
            "Формирует второй уровень нормативной иерархии.",
        ],
    )

    doc.add_heading("3.2 Контур сварного соединения и схем контроля", level=2)
    doc.add_paragraph("Таблица type_of_welded_joint")
    add_bullets(
        doc,
        [
            "Хранит типы сварных соединений.",
            "PK: id.",
            "FK: regulatory_documents_id -> regulatory_documents.id.",
            "Поля name и image_ref позволяют хранить как текстовый, так и графический контекст.",
        ],
    )
    doc.add_paragraph("Таблица params_by_type_welding_joint")
    add_bullets(
        doc,
        [
            "Содержит динамические параметры, добавляемые в карту в зависимости от типа шва.",
            "PK: id.",
            "FK: type_of_welded_joint_id -> type_of_welded_joint.id.",
            "Поля name/subtitle поддерживают UI-группировку параметров.",
        ],
    )
    doc.add_paragraph("Таблица cheme_control")
    add_bullets(
        doc,
        [
            "Справочник схем контроля (название и ссылка на изображение).",
            "PK: id.",
            "Используется через промежуточную таблицу связей.",
        ],
    )
    doc.add_paragraph("Таблица weld_type_to_scheme")
    add_bullets(
        doc,
        [
            "Реализует отношение many-to-many между типом шва и схемой контроля.",
            "PK: id.",
            "FK: type_of_welded_joint_id -> type_of_welded_joint.id.",
            "FK: cheme_control_id -> cheme_control.id.",
            "Ключевой механизм адаптации возможных схем к конкретному типу шва.",
        ],
    )
    doc.add_paragraph("Таблица width_of_controlled_area")
    add_bullets(
        doc,
        [
            "Содержит правила выбора ширины контролируемой зоны по диапазонам диаметра.",
            "PK: id.",
            "FK: type_of_welded_joint_id -> type_of_welded_joint.id.",
            "Пара (min_nominal_diametr, max_nominal_diam) описывает интервал применимости.",
        ],
    )

    doc.add_heading("3.3 Контур материалов, источников и режимов", level=2)
    doc.add_paragraph("Таблица type_metall")
    add_bullets(
        doc,
        [
            "Справочник материалов (железо, алюминий, магний, титан, медь, никель и др.).",
            "PK: id.",
            "Поле id_standard связывает материал с нормативным стандартом чувствительности.",
        ],
    )
    doc.add_paragraph("Таблица voltage_tube")
    add_bullets(
        doc,
        [
            "Таблица технологического ограничения напряжения трубки.",
            "PK: id.",
            "FK: type_metall_id -> type_metall.id.",
            "Пара (radiation_thickness, voltage_tube) задает пороговое правило выбора.",
            "Используется при фильтрации рентгеновских аппаратов и заполнении поля «Напряжение ... не более».",
        ],
    )
    doc.add_paragraph("Таблица rengen_apparatus")
    add_bullets(
        doc,
        [
            "Справочник рентгеновских аппаратов.",
            "PK: id.",
            "Поля: name, val (display), voltage_on_tube, focal_spot_size.",
            "В прикладной логике аппарат допустим только при соблюдении ограничений по напряжению.",
        ],
    )
    doc.add_paragraph("Таблица class_film_radiation_source")
    add_bullets(
        doc,
        [
            "Технологическая таблица выбора класса пленки по материалу и диапазону радиационной толщины.",
            "PK: id.",
            "FK: type_metall_id -> type_metall.id.",
            "Поля class_a/class_b/class_c отражают ветвление по уровню качества контроля.",
        ],
    )
    doc.add_paragraph("Таблица radiographic_film")
    add_bullets(
        doc,
        [
            "Справочник марок/типов радиографической пленки с указанием film_class.",
            "PK: id.",
            "Связывается с логикой выбора пленки по классу.",
        ],
    )
    doc.add_paragraph("Таблица control_sensitivity")
    add_bullets(
        doc,
        [
            "Таблица чувствительности контроля по диапазону толщин.",
            "PK: id.",
            "Поля class_a/class_b/class_c обеспечивают параметризацию расчета по уровню качества.",
        ],
    )

    doc.add_heading("3.4 Контур операций и шаблонов", level=2)
    doc.add_paragraph("Таблицы list_operation и operation_param")
    add_bullets(
        doc,
        [
            "list_operation хранит именованные наборы операций.",
            "operation_param содержит элементы набора (name_param, val, val2).",
            "FK: operation_param.id_list -> list_operation.id.",
            "Составной PK в operation_param (id, id_list) указывает, что id_list входит в идентичность записи.",
            "Данные используются при сборке блока «ПЕРЕЧЕНЬ ОПЕРАЦИЙ РК».",
        ],
    )

    doc.add_heading("3.5 Контур сохранения пользовательских карт", level=2)
    doc.add_paragraph("Таблица tech_cards")
    add_bullets(
        doc,
        [
            "PK: id (bigserial).",
            "Поле name — человекочитаемое имя снапшота.",
            "Поле card_data JSONB — полное состояние техкарты.",
            "created_at и updated_at — аудит жизненного цикла записи.",
        ],
    )
    doc.add_paragraph(
        "Использование JSONB здесь особенно важно: структура техкарты может эволюционировать "
        "без жестких миграций схемы на каждый новый параметр. Это стратегически снижает "
        "стоимость изменения фронтенда и серверных правил."
    )

    doc.add_heading("4. Кардинальности и связь сущностей", level=1)
    add_numbers(
        doc,
        [
            "control_methods (1) -> (N) regulatory_documents.",
            "regulatory_documents (1) -> (N) type_of_welded_joint.",
            "type_of_welded_joint (1) -> (N) params_by_type_welding_joint.",
            "type_of_welded_joint (N) <-> (N) cheme_control через weld_type_to_scheme.",
            "type_of_welded_joint (1) -> (N) width_of_controlled_area.",
            "type_metall (1) -> (N) class_film_radiation_source.",
            "type_metall (1) -> (N) voltage_tube.",
            "list_operation (1) -> (N) operation_param.",
        ],
    )
    doc.add_paragraph(
        "Таким образом, модель сочетает иерархические связи, табличные правила и "
        "ассоциативные связи many-to-many, что естественно для сложной инженерной "
        "предметной области."
    )

    doc.add_heading("5. Нормализация и целостность", level=1)
    doc.add_paragraph(
        "Схема преимущественно нормализована до 3НФ: справочные данные разделены по темам, "
        "повторяющиеся значения вынесены в отдельные таблицы, а связи представлены через FK. "
        "Это минимизирует дублирование и снижает риск аномалий обновления."
    )
    add_bullets(
        doc,
        [
            "Первичные ключи заданы почти во всех таблицах через serial/bigserial.",
            "Внешние ключи обеспечивают референциальную целостность без каскадного удаления.",
            "Отсутствие ON DELETE CASCADE оправдано: нормативные данные нельзя удалять неявно.",
            "JSONB в tech_cards сознательно вводит ограниченную денормализацию как компромисс гибкости.",
        ],
    )

    doc.add_heading("6. Интересные и сильные моменты модели", level=1)
    add_bullets(
        doc,
        [
            "Комбинация «справочник + табличное правило»: type_metall + voltage_tube.",
            "Ассоциативная таблица weld_type_to_scheme позволяет гибко расширять допустимые схемы без изменения кода.",
            "params_by_type_welding_joint делает UI динамическим: набор полей определяется данными, а не только шаблоном.",
            "operation_param поддерживает шаблоны операций и их централизованную настройку.",
            "tech_cards(JSONB) решает задачу версионирования и сохранения полного снимка пользовательского состояния.",
            "Разделение class_film_radiation_source и radiographic_film дает двухступенчатую модель выбора пленки: сначала класс, затем конкретная марка.",
        ],
    )

    doc.add_heading("7. Потоки использования данных в приложении", level=1)
    add_numbers(
        doc,
        [
            "Пользователь выбирает метод контроля -> подгружаются нормативные документы.",
            "Выбор документа определяет типы сварных соединений.",
            "Выбор типа шва определяет схемы контроля и набор специфических параметров.",
            "Материал и толщина определяют максимально допустимое напряжение трубки через voltage_tube.",
            "По допустимому напряжению фильтруется перечень ИИИ из rengen_apparatus.",
            "Класс качества и толщина влияют на выбор чувствительности и пленки.",
            "Итоговая карта сохраняется в tech_cards как снапшот.",
        ],
    )

    doc.add_heading("8. Ограничения текущей схемы и точки роста", level=1)
    add_bullets(
        doc,
        [
            "Во многих таблицах полезны дополнительные UNIQUE-ограничения (например, material в type_metall, name в справочниках).",
            "Для производительности целесообразны индексы по часто фильтруемым полям: lower(trim(material)), type_metall_id + radiation_thickness.",
            "В operation_param составной PK можно дополнить бизнес-ограничениями порядка/уникальности name_param внутри списка.",
            "Для аудита изменений нормативных таблиц полезно ввести версионирование (valid_from/valid_to).",
            "Для tech_cards при масштабировании могут потребоваться GIN-индексы по card_data для аналитических выборок.",
        ],
    )

    doc.add_heading("9. Рекомендуемые SQL-улучшения (дипломный раздел)", level=1)
    doc.add_paragraph(
        "Ниже приведены практические направления улучшения схемы, которые можно включить "
        "в раздел «Проектирование и оптимизация БД» дипломной работы:"
    )
    add_numbers(
        doc,
        [
            "Добавить UNIQUE(material) в type_metall для устранения дублирования материалов.",
            "Добавить индекс idx_voltage_tube_lookup(type_metall_id, radiation_thickness).",
            "Добавить UNIQUE(type_of_welded_joint_id, cheme_control_id) в weld_type_to_scheme.",
            "Добавить CHECK-ограничения неотрицательности для толщин и напряжений.",
            "Ввести таблицу версий нормативных данных для прослеживаемости изменений.",
            "Внедрить триггер автоматического обновления updated_at в tech_cards.",
        ],
    )

    doc.add_heading("10. Заключение", level=1)
    doc.add_paragraph(
        "Спроектированная база данных демонстрирует удачный баланс между строгой "
        "реляционной структурой нормативных данных и гибкостью хранения пользовательского "
        "состояния в JSONB. Модель хорошо соответствует инженерной природе задачи: "
        "критические зависимости формализованы внешними ключами, а вычислительные правила "
        "выведены в отдельные таблицы, которые можно актуализировать без радикальной "
        "переработки прикладного кода. В контексте дипломной работы это является важным "
        "аргументом в пользу масштабируемости, сопровождаемости и практической применимости решения."
    )

    doc.save(OUT_PATH)
    print(OUT_PATH)


if __name__ == "__main__":
    main()
