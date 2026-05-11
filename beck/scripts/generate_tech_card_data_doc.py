"""
Генерация Word-документа: класс TechCardData, описание методов, листинги из services/tech_card.py.

Запуск: python scripts/generate_tech_card_data_doc.py
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1]
SOURCE_FILE = ROOT / "services" / "tech_card.py"
OUTPUT = ROOT / "docs" / "tech_card_data_reference.docx"

# Описания методов (дипломный стиль; синхронизировать при изменении API)
METHOD_DESCRIPTIONS: list[tuple[str, str]] = [
    ("__init__", "Инициализация контейнера карты: словарь блоков `params` (ключ верхнего уровня — id блока)."),
    ("get", "Возвращает блок по строковому ключу верхнего уровня (`params[key]`)."),
    ("set", "Записывает блок по ключу верхнего уровня."),
    ("to_dict", "Экспорт для совместимости: `{\"currentParams\": self.params}`."),
    ("_to_json_dict", "Внутреннее представление для JSON: `{\"params\": self.params}`."),
    ("__str__", "Строковое представление — сериализованный JSON с отступами."),
    ("serialise", "Сериализация карты в JSON-строку (UTF-8, без ASCII-экранирования)."),
    (
        "from_jsonDeSerialise",
        "Десериализация из строки или словаря: извлекает `type`, `methodology`, `params` и заполняет объект.",
    ),
    (
        "has_block_and_param",
        "Проверяет, существует ли блок с данным `name` и параметр с данным именем внутри него.",
    ),
    ("has_block", "Проверяет наличие блока по полю `name` среди значений `params`."),
    (
        "_find_free_id",
        "Подбирает минимальный свободный целочисленный ключ среди ключей словаря параметров (int или строки-цифры).",
    ),
    (
        "add_param_to_block",
        "Добавляет параметр в конец блока с новым свободным id; запрещает дубликат по `name`; требует ключ `name` в словаре параметра.",
    ),
    (
        "insert_param_to_block",
        "Вставляет параметр на позицию `insert_id` со сдвигом занятых целочисленных ключей вправо; запрещает дубликат по имени.",
    ),
    (
        "_id_to_sort_key",
        "Преобразует ключ параметра (int или строка вида «1.2.10») в кортеж для лексикографической сортировки порядка полей.",
    ),
    ("sort_all_params", "Сортирует параметры в каждом блоке по ключам через `_id_to_sort_key`."),
    ("get_param_value", "Возвращает `val` первого параметра с заданным именем в блоке с заданным именем."),
    ("set_param_value", "Устанавливает `val` у параметра с заданным именем в блоке."),
    (
        "update_param",
        "Частичное обновление словаря параметра (`dict.update`): options, subtitle, метаданные и т.д.",
    ),
    (
        "remove_param_from_block",
        "Удаляет первый параметр с указанным именем из блока (по совпадению ключа в словаре `params` блока).",
    ),
    ("_parse_id", "Разбор ключа параметра в список целых компонент (иерархические id)."),
    ("_same_level", "Сравнение глубины иерархии двух id (одинаковая длина списка компонент)."),
    (
        "change_param_id_by_name_autoshift",
        "Перенос параметра на новый иерархический id со сдвигом элементов того же уровня; затем сортировка.",
    ),
    (
        "hasSpecParam",
        "Возвращает False, если `val` отсутствует или является списком/кортежем (каталог без выбора); иначе True.",
    ),
    (
        "insert_param_to_block_reWrite",
        "Вставка/замена параметра по составному `insert_id` (строка «4.2» и т.п.) со сдвигом ключей; при совпадении имени старый удаляется.",
    ),
    ("_key_to_tuple", "Преобразование ключа словаря параметров в tuple целых."),
    ("_tuple_to_key", "Обратное преобразование: tuple → строковый ключ с точками."),
    ("_compare_tuples", "Лексикографическое сравнение tuple-ключей."),
    ("_increment_tuple", "Инкремент компонента ключа при сдвиге при вставке (вспомогательная для reWrite)."),
    ("_decrement_tuple", "Декремент целевой позиции вставки при удалении предшествующего ключа."),
]


def _set_body(doc: Document) -> None:
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(14)
    pf = style.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_after = Pt(6)


def _add_code_lines(doc: Document, text: str, font_pt: float = 9) -> None:
    """Моноширинный листинг: по одной строке на абзац (удобно для диплома и копирования)."""
    for line in text.splitlines():
        p = doc.add_paragraph()
        run = p.add_run(line if line else " ")
        run.font.name = "Courier New"
        run.font.size = Pt(font_pt)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        fmt = p.paragraph_format
        fmt.left_indent = Pt(28)


def _add_method_table(doc: Document) -> None:
    table = doc.add_table(rows=1, cols=2)
    hdr = table.rows[0].cells
    hdr[0].text = "Метод"
    hdr[1].text = "Назначение"
    for name, desc in METHOD_DESCRIPTIONS:
        row = table.add_row().cells
        row[0].text = name
        row[1].text = desc


def main() -> None:
    if not SOURCE_FILE.is_file():
        raise SystemExit(f"Не найден файл: {SOURCE_FILE}")

    source_text = SOURCE_FILE.read_text(encoding="utf-8")
    lines = source_text.splitlines()

    doc = Document()
    _set_body(doc)

    doc.add_heading(
        "Класс TechCardData: структура данных операционной технологической карты",
        0,
    )
    doc.add_paragraph(
        f"Документ подготовлен для пояснительной записки. Исходный модуль: "
        f"`services/tech_card.py`. Дата генерации: {datetime.now().strftime('%d.%m.%Y %H:%M')}."
    )
    doc.add_paragraph(
        "TechCardData — центральная модель документа «технологическая карта»: иерархия блоков и параметров "
        "хранится во вложенных словарях, что соответствует формату JSON при обмене с клиентом и сохранении снимков."
    )

    doc.add_heading("1. Импорты, перечисление TypeObjectControl и сериализация", level=1)
    doc.add_paragraph(
        "В модуле объявлено перечисление `StrEnum` для типа объекта контроля (пластина / труба). "
        "Ниже — листинг от начала файла до методов поиска блоков (включая конструктор, get/set и JSON)."
    )
    frag1_end = next(
        (i for i, line in enumerate(lines) if line.startswith("    def has_block_and_param")),
        len(lines),
    )
    _add_code_lines(doc, "\n".join(lines[:frag1_end]))

    doc.add_heading("2. Сводная таблица методов класса TechCardData", level=1)
    doc.add_paragraph(
        "В таблице приведены открытые и служебные методы, используемые сервисом и changers. "
        "Названия частных методов (с префиксом «_») отражают внутреннюю механику сортировки и иерархических ключей."
    )
    _add_method_table(doc)

    doc.add_heading("3. Листинг: доступ, вставка и сортировка параметров", level=1)
    doc.add_paragraph(
        "Фрагмент: проверки наличия блоков, добавление и вставка параметров, сортировка по ключам."
    )
    start = next((i for i, l in enumerate(lines) if l.startswith("    def has_block_and_param")), 0)
    end = next((i for i, l in enumerate(lines) if l.startswith("    def get_param_value")), len(lines))
    _add_code_lines(doc, "\n".join(lines[start:end]))

    doc.add_heading("4. Листинг: чтение, запись, обновление, удаление", level=1)
    start = next((i for i, l in enumerate(lines) if l.startswith("    def get_param_value")), 0)
    end = next((i for i, l in enumerate(lines) if l.startswith("    def _parse_id")), len(lines))
    _add_code_lines(doc, "\n".join(lines[start:end]))

    doc.add_heading("5. Листинг: иерархические идентификаторы и перенос параметра", level=1)
    start = next((i for i, l in enumerate(lines) if l.startswith("    def _parse_id")), 0)
    end = next((i for i, l in enumerate(lines) if l.startswith("    def hasSpecParam")), len(lines))
    _add_code_lines(doc, "\n".join(lines[start:end]))

    doc.add_heading("6. Листинг: hasSpecParam и insert_param_to_block_reWrite", level=1)
    start = next((i for i, l in enumerate(lines) if l.startswith("    def hasSpecParam")), 0)
    _add_code_lines(doc, "\n".join(lines[start:]))

    doc.add_heading("Приложение. Полный листинг модуля tech_card.py", level=1)
    doc.add_paragraph(
        "Ниже приведён полный текст файла на момент генерации документа (удобно для приложения к ВКР)."
    )
    _add_code_lines(doc, source_text, font_pt=8)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(str(OUTPUT))


if __name__ == "__main__":
    main()
