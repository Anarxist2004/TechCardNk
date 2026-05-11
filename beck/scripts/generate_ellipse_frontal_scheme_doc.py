from __future__ import annotations

from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.shared import Pt


ROOT = Path(r"e:\sem8\ful\Diplom_NK\beck")
OUTPUT = ROOT / "docs" / "rengen_distance_for_ellipse_frontal_scheme_detailed.docx"


def _setup(doc: Document) -> None:
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(14)
    pf = style.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_after = Pt(6)


def _p(doc: Document, text: str, bullet: bool = False) -> None:
    doc.add_paragraph(text, style="List Bullet" if bullet else None)


def _h(doc: Document, text: str, level: int) -> None:
    doc.add_heading(text, level=level)


def main() -> None:
    doc = Document()
    _setup(doc)

    doc.add_heading(
        "Подробное описание changer RengenDistanceForEllipseFrontalScheme",
        0,
    )
    _p(
        doc,
        f"Документ подготовлен по модулю "
        f"`services/Changers/ch_RengenDistanceForEllipseFrontalScheme.py`. "
        f"Дата: {datetime.now().strftime('%d.%m.%Y %H:%M')}.",
    )
    _p(
        doc,
        "Назначение changer-а: для схемы 7.3 "
        "«фронтальное просвечивание через две стенки на эллипс» "
        "рассчитать нижнюю границу расстояния f от источника излучения до "
        "контролируемого соединения и вывести инженерную подсказку в техкарте.",
    )

    _h(doc, "1. Входные параметры и обозначения", 1)
    for line in [
        "D — номинальный наружный диаметр трубы, мм.",
        "S_st — номинальная толщина стенки, мм.",
        "d — внутренний диаметр трубы, мм.",
        "K — чувствительность контроля, мм.",
        "Φ — размер (максимум) фокусного пятна ИИИ, мм.",
        "Q — уровень качества изображения (A, B, C).",
        "S_q — коэффициент класса изображения (A->1.2, B->1.1, C->1.5).",
        "f_min — минимально допустимое расстояние от ИИИ до поверхности.",
    ]:
        _p(doc, line, bullet=True)

    _h(doc, "2. Условия активации и проверки применимости", 1)
    _p(
        doc,
        "Расчет включается только для схемы, определяемой по признакам "
        "«фронтальное просвечивание» и «на эллипс» в значении параметра "
        "«схема просвечивания».",
    )
    for line in [
        "Требуется наличие блока «ИСХОДНЫЕ ДАННЫЕ».",
        "Обязательны валидные D, S_st, K, Q.",
        "Обязателен факт выбранного аппарата (scalar-значение в поле «ИИИ»).",
        "Обязателен корректно распарсенный Φ из поля фокусного пятна.",
        "Геометрия должна быть корректной: d > 0 и D > d.",
    ]:
        _p(doc, line, bullet=True)

    _h(doc, "3. Формулы и инженерная зависимость", 1)
    _h(doc, "3.1 Геометрия", 2)
    _p(doc, "Внутренний диаметр:")
    _p(doc, "d = D - 2*S_st", bullet=True)

    _h(doc, "3.2 Коэффициент по классу качества", 2)
    _p(doc, "Для уровня качества выбирается множитель S_q:")
    _p(doc, "Q = A -> S_q = 1.2", bullet=True)
    _p(doc, "Q = B -> S_q = 1.1", bullet=True)
    _p(doc, "Q = C -> S_q = 1.5", bullet=True)

    _h(doc, "3.3 Основная расчетная формула", 2)
    _p(doc, "В модуле реализовано условие:")
    _p(doc, "f >= (2 * Φ * d * S_q) / K", bullet=True)
    _p(doc, "Т.е. минимальная граница:")
    _p(doc, "f_min = (2 * Φ * d * S_q) / K", bullet=True)
    _p(doc, "После вычисления выполняется отсечение снизу:")
    _p(doc, "если f_min < 0, то f_min := 0", bullet=True)
    _p(
        doc,
        "В карту записывается подсказка вида «f>=X», где X — форматированное "
        "значение f_min.",
    )

    _h(doc, "4. Пошаговый алгоритм работы changer-а", 1)
    for step in [
        "Проверка наличия блока «ИСХОДНЫЕ ДАННЫЕ».",
        "Проверка, что выбрана именно схема 7.3 «на эллипс».",
        "Считывание и парсинг D, S_st, K.",
        "Считывание и интерпретация уровня качества Q.",
        "Проверка, что аппарат ИИИ выбран пользователем (не каталог, а scalar).",
        "Чтение и парсинг фокусного пятна Φ.",
        "Расчет d = D - 2*S_st и геометрическая валидация.",
        "Определение S_q по Q.",
        "Расчет f_min = (2*Φ*d*S_q)/K.",
        "Отсечение f_min до 0 при необходимости.",
        "Запись подсказки «f>=...».",
    ]:
        _p(doc, step, bullet=True)

    _h(doc, "5. Поведение при неполных или некорректных данных", 1)
    _p(
        doc,
        "Алгоритм построен по принципу безопасной деградации: при недостаточности "
        "исходных данных числовая рекомендация не формируется.",
    )
    for line in [
        "Если схема не 7.3 — changer не изменяет расстояние.",
        "Если отсутствует выбранный аппарат — расстояние не рассчитывается.",
        "Если не удалось распарсить фокус Φ — расчет не выполняется.",
        "Если K<=0, D<=0 или геометрия d<=0 / D<=d — расчет не выполняется.",
        "Во всех таких случаях placeholder/hint поля расстояния очищается (None).",
    ]:
        _p(doc, line, bullet=True)

    _h(doc, "6. Выходные данные в структуре техкарты", 1)
    for line in [
        "Обновляется поле расстояния от ИИИ до поверхности контролируемого соединения.",
        "Записываются сразу два атрибута: placeholder и hint (для совместимости фронта).",
        "Формат результата: «f>=X».",
    ]:
        _p(doc, line, bullet=True)

    _h(doc, "7. Вывод для ВКР", 1)
    _p(
        doc,
        "RengenDistanceForEllipseFrontalScheme реализует целевой инженерный расчет "
        "для схемы 7.3: минимальное расстояние определяется на основе геометрии трубы, "
        "параметров источника и требуемого класса качества. Модуль дополняет "
        "панорамно-геометрические расчеты системы и обеспечивает воспроизводимую "
        "формализацию нормативного правила в вычислительном виде.",
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(str(OUTPUT))


if __name__ == "__main__":
    main()
