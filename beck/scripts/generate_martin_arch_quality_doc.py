from __future__ import annotations

import ast
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Set, Tuple

from docx import Document
from docx.shared import Pt


ROOT = Path(r"e:\sem8\ful\Diplom_NK\beck")
TARGET_LAYERS = ("adapter", "controllers", "services", "repositories")
# Backward-compatible alias for Excel script imports.
TARGET_PACKAGES = TARGET_LAYERS
OUTPUT = ROOT / "docs" / "architecture_quality_martin.docx"
_RUNTIME_ROOT_PACKAGES = ("controllers", "services", "repositories")
_EXCLUDED_PARTS = {"test", "tests", "__pycache__", "docs", "res", "assets", "scripts"}


@dataclass
class ModuleInfo:
    module: str
    package: str
    layer: str
    path: Path
    imports: Set[str] = field(default_factory=set)
    classes: int = 0
    abstract_classes: int = 0


def _module_name(py_file: Path) -> str:
    rel = py_file.relative_to(ROOT).as_posix()
    return rel[:-3].replace("/", ".")


def _extract_imports(tree: ast.AST, current_module: str) -> Set[str]:
    out: Set[str] = set()
    parts = current_module.split(".")
    for node in ast.walk(tree):
        if isinstance(node, ast.Import):
            for alias in node.names:
                out.add(alias.name)
        elif isinstance(node, ast.ImportFrom):
            if node.level and node.module is not None:
                base_parts = parts[:-node.level]
                out.add(".".join(base_parts + node.module.split(".")))
            elif node.level and node.module is None:
                base_parts = parts[:-node.level]
                if base_parts:
                    out.add(".".join(base_parts))
            elif node.module:
                out.add(node.module)
    return {x for x in out if x}


def _is_abstract_class(node: ast.ClassDef) -> bool:
    for base in node.bases:
        if isinstance(base, ast.Name) and base.id == "ABC":
            return True
        if isinstance(base, ast.Attribute) and base.attr == "ABC":
            return True
    for item in node.body:
        if isinstance(item, (ast.FunctionDef, ast.AsyncFunctionDef)):
            for dec in item.decorator_list:
                if isinstance(dec, ast.Name) and dec.id == "abstractmethod":
                    return True
                if isinstance(dec, ast.Attribute) and dec.attr == "abstractmethod":
                    return True
    return False


def collect_modules() -> Dict[str, ModuleInfo]:
    modules: Dict[str, ModuleInfo] = {}
    for py in ROOT.rglob("*.py"):
        parts_lower = [p.lower() for p in py.parts]
        if any(p in _EXCLUDED_PARTS for p in parts_lower):
            continue
        if any(p.startswith("test") for p in parts_lower):
            continue
        module = _module_name(py)
        package = module.split(".")[0]
        if package not in _RUNTIME_ROOT_PACKAGES:
            continue
        layer = "adapter" if module.startswith("controllers.adapter") else package

        text = py.read_text(encoding="utf-8")
        tree = ast.parse(text)
        imports = _extract_imports(tree, module)
        classes = 0
        abstract_classes = 0
        for node in ast.walk(tree):
            if isinstance(node, ast.ClassDef):
                classes += 1
                if _is_abstract_class(node):
                    abstract_classes += 1

        modules[module] = ModuleInfo(
            module=module,
            package=package,
            layer=layer,
            path=py,
            imports=imports,
            classes=classes,
            abstract_classes=abstract_classes,
        )
    return modules


def build_dependency_sets(
    modules: Dict[str, ModuleInfo],
) -> Tuple[Dict[str, Set[str]], Dict[str, Set[str]]]:
    internal_modules = set(modules.keys())
    imports_by_module: Dict[str, Set[str]] = {m: set() for m in internal_modules}

    for m, info in modules.items():
        for imp in info.imports:
            for target in internal_modules:
                if imp == target or imp.startswith(target + "."):
                    imports_by_module[m].add(target)
                    break

    ce: Dict[str, Set[str]] = {p: set() for p in TARGET_LAYERS}
    ca: Dict[str, Set[str]] = {p: set() for p in TARGET_LAYERS}

    for src, targets in imports_by_module.items():
        src_layer = modules[src].layer
        for dst in targets:
            dst_layer = modules[dst].layer
            if src_layer != dst_layer:
                ce[src_layer].add(dst)
                ca[dst_layer].add(src)

    return ce, ca


def count_package_classes(modules: Dict[str, ModuleInfo]) -> Dict[str, Tuple[int, int]]:
    out: Dict[str, Tuple[int, int]] = {}
    for layer in TARGET_LAYERS:
        cls = sum(m.classes for m in modules.values() if m.layer == layer)
        abs_cls = sum(m.abstract_classes for m in modules.values() if m.layer == layer)
        out[layer] = (cls, abs_cls)
    return out


def compute_metrics(modules: Dict[str, ModuleInfo]):
    ce, ca = build_dependency_sets(modules)
    class_counts = count_package_classes(modules)
    metrics = {}
    for layer in TARGET_LAYERS:
        nc, na = class_counts[layer]
        ce_n = len(ce[layer])
        ca_n = len(ca[layer])
        a = (na / nc) if nc else 0.0
        i = (ce_n / (ca_n + ce_n)) if (ca_n + ce_n) else 0.0
        d = abs(a + i - 1.0)
        metrics[layer] = {
            "Nc": nc,
            "Na": na,
            "Ca": ca_n,
            "Ce": ce_n,
            "A": a,
            "I": i,
            "D": d,
        }
    return metrics


def layer_violations(modules: Dict[str, ModuleInfo]) -> List[str]:
    # Пользовательская модель слоев:
    # adapter -> controllers -> services -> repositories
    # Прямые "перепрыгивания" слоев считаем нарушением.
    violations: List[str] = []
    module_names = set(modules)
    allowed: Dict[str, Set[str]] = {
        "adapter": {"adapter", "controllers"},
        "controllers": {"controllers", "services"},
        "services": {"services", "repositories"},
        "repositories": {"repositories"},
    }
    for src, info in modules.items():
        src_layer = info.layer
        for imp in info.imports:
            for target in module_names:
                if not (imp == target or imp.startswith(target + ".")):
                    continue
                dst_layer = modules[target].layer
                if src_layer == dst_layer:
                    continue
                if dst_layer not in allowed[src_layer]:
                    violations.append(f"{src} -> {target}")
                break
    return sorted(set(violations))


def package_cycles(modules: Dict[str, ModuleInfo]) -> List[str]:
    edges: Dict[str, Set[str]] = {p: set() for p in TARGET_LAYERS}
    module_names = set(modules)
    for src, info in modules.items():
        src_layer = info.layer
        for imp in info.imports:
            for target in module_names:
                if imp == target or imp.startswith(target + "."):
                    dst_layer = modules[target].layer
                    if dst_layer != src_layer:
                        edges[src_layer].add(dst_layer)
                    break

    cycles: List[str] = []
    for a in TARGET_LAYERS:
        for b in edges[a]:
            if a in edges[b]:
                pair = " <-> ".join(sorted((a, b)))
                if pair not in cycles:
                    cycles.append(pair)
    return sorted(cycles)


def score(metrics: Dict[str, dict], violations: List[str], cycles: List[str]) -> Dict[str, float]:
    avg_d = sum(v["D"] for v in metrics.values()) / len(metrics)
    main_seq_score = (1.0 - avg_d) * 100.0

    violation_penalty = min(100.0, len(violations) * 5.0)
    layer_score = max(0.0, 100.0 - violation_penalty)

    cycle_penalty = min(100.0, len(cycles) * 25.0)
    cycle_score = max(0.0, 100.0 - cycle_penalty)

    total = 0.5 * main_seq_score + 0.3 * layer_score + 0.2 * cycle_score
    return {
        "main_seq_score": main_seq_score,
        "layer_score": layer_score,
        "cycle_score": cycle_score,
        "total": total,
    }


def fmt(x: float) -> str:
    return f"{x:.3f}"


def build_doc(metrics: Dict[str, dict], violations: List[str], cycles: List[str], scores: Dict[str, float], modules: Dict[str, ModuleInfo]) -> None:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    doc.add_heading("Расчет качества архитектуры по Роберту Мартину", 0)
    doc.add_paragraph(f"Дата: {datetime.now().strftime('%d.%m.%Y %H:%M')}")
    doc.add_paragraph(
        "Документ рассчитывает ключевые метрики из подхода Роберта Мартина для текущего "
        "состояния проекта: Abstractness (A), Instability (I), Distance from Main Sequence (D), "
        "а также проверяет направленность зависимостей и наличие циклов между слоями."
    )

    doc.add_heading("1. Формулы", level=1)
    for line in [
        "A = Na / Nc, где Na — число абстрактных классов, Nc — общее число классов в пакете.",
        "I = Ce / (Ca + Ce), где Ce — исходящие зависимости пакета, Ca — входящие.",
        "D = |A + I - 1| — расстояние до «главной последовательности».",
        "Чем ближе D к 0, тем более сбалансирован пакет.",
    ]:
        doc.add_paragraph(line, style="List Bullet")

    doc.add_heading("2. Область расчета", level=1)
    doc.add_paragraph("Анализированы слои: adapter, controllers, services, repositories.")
    doc.add_paragraph(
        "Из расчета исключены тестовые/вспомогательные модули (test/tests), "
        "а также нерелевантные runtime-директории ресурсов."
    )
    doc.add_paragraph(f"Количество runtime-модулей в расчете: {len(modules)}.")

    doc.add_heading("3. Метрики по слоям", level=1)
    table = doc.add_table(rows=1, cols=8)
    hdr = table.rows[0].cells
    hdr[0].text = "Слой"
    hdr[1].text = "Nc"
    hdr[2].text = "Na"
    hdr[3].text = "Ca"
    hdr[4].text = "Ce"
    hdr[5].text = "A"
    hdr[6].text = "I"
    hdr[7].text = "D"
    for pkg in TARGET_LAYERS:
        row = table.add_row().cells
        m = metrics[pkg]
        row[0].text = pkg
        row[1].text = str(m["Nc"])
        row[2].text = str(m["Na"])
        row[3].text = str(m["Ca"])
        row[4].text = str(m["Ce"])
        row[5].text = fmt(m["A"])
        row[6].text = fmt(m["I"])
        row[7].text = fmt(m["D"])

    doc.add_heading("4. Проверка правил направленности зависимостей", level=1)
    if violations:
        doc.add_paragraph("Найдены нарушения (внутренние зависимости слоев):")
        for v in violations:
            doc.add_paragraph(v, style="List Bullet")
    else:
        doc.add_paragraph("Нарушений не обнаружено.")

    doc.add_heading("5. Проверка циклов между слоями", level=1)
    if cycles:
        doc.add_paragraph("Обнаружены циклы на уровне слоев:")
        for c in cycles:
            doc.add_paragraph(c, style="List Bullet")
    else:
        doc.add_paragraph("Циклических зависимостей между слоями не найдено.")

    doc.add_heading("6. Интегральный индекс качества (рабочая шкала)", level=1)
    doc.add_paragraph(
        "Индекс составлен как инженерная агрегатная оценка:\n"
        "50% — близость к Main Sequence,\n"
        "30% — соблюдение направленности слоев,\n"
        "20% — отсутствие циклов."
    )
    doc.add_paragraph(f"Main Sequence Score: {scores['main_seq_score']:.2f} / 100")
    doc.add_paragraph(f"Layering Score: {scores['layer_score']:.2f} / 100")
    doc.add_paragraph(f"Cycle Score: {scores['cycle_score']:.2f} / 100")
    doc.add_paragraph(f"Итоговый Architecture Quality Index: {scores['total']:.2f} / 100")

    doc.add_heading("7. Интерпретация и рекомендации", level=1)
    doc.add_paragraph(
        "D близкое к 0 говорит о балансе между абстрактностью и стабильностью. "
        "Высокий D у стабильного пакета при низкой абстрактности обычно указывает на "
        "технический долг. Нарушения направленности и циклы ухудшают сопровождаемость "
        "и тестируемость, особенно при росте проекта."
    )
    doc.add_paragraph(
        "Рекомендуемые приоритеты: устранение циклов между слоями, снижение зависимостей "
        "repositories от services, декомпозиция крупных пакетов и выделение контрактов "
        "в интерфейсы."
    )

    doc.save(OUTPUT)


def main() -> None:
    modules = collect_modules()
    metrics = compute_metrics(modules)
    violations = layer_violations(modules)
    cycles = package_cycles(modules)
    scores = score(metrics, violations, cycles)
    build_doc(metrics, violations, cycles, scores, modules)
    print(str(OUTPUT))


if __name__ == "__main__":
    main()
