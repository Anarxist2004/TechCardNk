"""
Генерация Word-документа: веб-адаптер (единственный модуль controllers/adapterWeb.py).

В репозитории нет файла adapter.py — адаптер реализован в adapterWeb.py.

Запуск из корня beck: python scripts/generate_adapter_component_doc.py
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "adapter_component_detailed.docx"


def _set_body_style(doc: Document) -> None:
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    pf = style.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_after = Pt(6)


def _p(doc: Document, text: str) -> None:
    doc.add_paragraph(text)


def _heading(doc: Document, text: str, level: int) -> None:
    doc.add_heading(text, level=level)


def build_document() -> Document:
    doc = Document()
    _set_body_style(doc)

    doc.add_heading(
        "Компонент веб-адаптера (модуль adapterWeb.py)",
        0,
    )
    _p(
        doc,
        f"Документ сформирован по файлу `controllers/adapterWeb.py`. "
        f"Дата: {datetime.now().strftime('%d.%m.%Y %H:%M')}. "
        "В каталоге `controllers/` отдельного файла `adapter.py` нет: весь HTTP-вход "
        "приложения сосредоточен в указанном модуле.",
    )

    _heading(doc, "1. Роль компонента", 1)
    _p(
        doc,
        "Адаптер в смысле гексагональной («портов и адаптеров») архитектуры — это код на "
        "периферии системы, который переводит внешний протокол (здесь — HTTP JSON поверх "
        "FastAPI) во внутренние вызовы абстракции контроллера `IControllers`. Модуль "
        "не содержит бизнес-правил техкарты: он разбирает запрос, выбирает операцию по "
        "строковому параметру пути, передаёт данные в контроллер и возвращает результат "
        "клиенту либо формирует коды ошибок.",
    )

    _heading(doc, "2. Инициализация приложения FastAPI", 1)
    _p(
        doc,
        "На уровне модуля создаётся экземпляр `app = FastAPI()`. Вычисляется каталог "
        "`res` относительно корня проекта (`parent.parent` от файла адаптера); если "
        "каталог существует, монтируется раздача статики по префиксу `/res` через "
        "`StaticFiles` — это позволяет отдавать изображения и прочие ресурсы без "
        "отдельного веб-сервера для файлов.",
    )
    _p(
        doc,
        "Подключается middleware `CORSMiddleware` с широкими настройками для теста: "
        "`allow_origins=[\"*\"]`, разрешены все методы и заголовки, включены credentials. "
        "Для production-окружения такой профиль обычно ужесточают (конкретные origin’ы, "
        "ограниченный набор методов).",
    )

    _heading(doc, "3. Функция create_adapter", 1)
    _p(
        doc,
        "Сигнатура: `create_adapter(controller: IControllers)`. Внутри объявляется "
        "асинхронный обработчик `adapter` и регистрируется маршрут `POST /techcard/{control_type}`. "
        "Тип операции задаётся path-параметром `control_type`; тело запроса ожидается как "
        "JSON и попадает в аргумент `payload` через `Body(...)`. Если `payload` не словарь, "
        "используется пустой словарь.",
    )
    _p(
        doc,
        "Вспомогательная функция `safe_int` безопасно приводит значения к `int` с "
        "значением по умолчанию при ошибке; из тела извлекаются поля `methodology`, "
        "`idElement`, `type` (с дефолтами 0, 1, 1). Эти величины участвуют в логировании "
        "при сбоях, даже если для конкретной ветки `control_type` они не используются в "
        "вызове контроллера.",
    )

    _heading(doc, "4. Соответствие control_type и вызовов контроллера", 1)
    rows = [
        ("template", "`controller.get_template()` — без дополнительных полей из payload."),
        ("updateTechCard", "`controller.updateTechCard(request_payload.get(\"techCard\", {}))`."),
        (
            "saveTechCard",
            "`saveTechCard(name, data, id)` — имя из `name`, данные карты из `data`, "
            "необязательный идентификатор из `id`.",
        ),
        ("listSavedTechCards", "Список сохранённых карт без параметров из payload."),
        (
            "getSavedTechCard",
            "Идентификатор из `id` через `safe_int`; при `None` с сервиса — ответ 404 "
            "с текстом «Tech card not found».",
        ),
    ]
    for op, desc in rows:
        _p(doc, f"{op} — {desc}")
    _p(
        doc,
        "Любое другое значение `control_type` приводит к возврату пустого объекта `{}` "
        "с кодом 200 (исключение не генерируется).",
    )

    _heading(doc, "5. Успешный ответ и ошибки", 1)
    _p(
        doc,
        "После вызова контроллера результат печатается в консоль (`print`) и возвращается "
        "как тело ответа FastAPI (типично словарь или список, сериализуемый в JSON).",
    )
    _p(
        doc,
        "Любое необработанное исключение в блоке `try` перехватывается: в stdout уходит "
        "расширенное сообщение с `control_type`, methodology, idElement, type, списком "
        "ключей payload и текстом ошибки; выводится полный traceback; клиенту возвращается "
        "`HTTPException` со статусом 500 и detail с подсказкой смотреть лог бэкенда.",
    )

    _heading(doc, "6. Запуск сервера", 1)
    _p(
        doc,
        "В конце `create_adapter` вызывается `uvicorn.run(app, host=\"0.0.0.0\", port=8000)` — "
        "то есть при создании адаптера сразу блокирующий запуск ASGI-сервера на всех "
        "интерфейсах и порту 8000. Функция возвращает ссылку на вложенную функцию `adapter` "
        "(после `run` управление до `return` в обычном сценарии не доходит; `return` имеет "
        "смысл при иной организации тестов или если `uvicorn.run` заменён заглушкой).",
    )

    _heading(doc, "7. Связь с точкой входа", 1)
    _p(
        doc,
        "В `main.py` после сборки репозитория, pipeline и `TechCardService` создаётся "
        "`ControllerWeb(service)` и передаётся в `create_adapter(controller)`, что "
        "замыкает цепочку: HTTP → адаптер → контроллер → сервис → репозиторий.",
    )

    _heading(doc, "8. Замечания по исходнику", 1)
    _p(
        doc,
        "Импорты `Request` из `fastapi` и `BaseModel` из `pydantic` в текущей версии файла "
        "не используются. Переменная `urlObjec` и закомментированная строка в конце файла "
        "выглядят как задел под будущее расширение (например, `get_available_params_for_type`).",
    )

    return doc


def main() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = build_document()
    doc.save(OUTPUT)
    print(f"Written: {OUTPUT}")


if __name__ == "__main__":
    main()
