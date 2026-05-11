from datetime import datetime

from docx import Document
from docx.shared import Pt


OUT_PATH = r"e:\sem8\ful\Diplom_NK\beck\docs\database_chapter_full.docx"


def bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def nums(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Number")


def main() -> None:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    doc.add_heading("ГЛАВА 3. ПРОЕКТИРОВАНИЕ И РЕАЛИЗАЦИЯ БАЗЫ ДАННЫХ", 0)
    doc.add_paragraph(f"Сформировано автоматически: {datetime.now().strftime('%d.%m.%Y %H:%M')}")

    doc.add_heading("3.1 Цели и роль базы данных в системе", level=1)
    doc.add_paragraph(
        "В рамках разрабатываемой системы радиографического контроля база данных выполняет "
        "функции нормативного ядра и хранилища эксплуатационных данных. База данных обеспечивает "
        "централизованное управление справочниками, таблицами инженерных зависимостей и "
        "состояниями сформированных технологических карт."
    )
    doc.add_paragraph(
        "Ключевая задача БД — предоставить приложению надежный и формализованный источник "
        "данных, позволяющий автоматически рассчитывать и подбирать параметры контроля по "
        "входным характеристикам объекта (тип шва, материал, толщина, схема просвечивания и т.д.)."
    )

    doc.add_heading("3.2 Требования к данным и ограничения предметной области", level=1)
    bullets(
        doc,
        [
            "Должна поддерживаться иерархия: метод контроля -> нормативный документ -> тип сварного соединения.",
            "Для каждого типа шва должны храниться динамические параметры и допустимые схемы контроля.",
            "Требуется хранение правил выбора напряжения трубки по материалу и радиационной толщине.",
            "Должны поддерживаться справочники источников излучения, пленок, чувствительности и экрана.",
            "Система должна сохранять полное состояние пользовательской техкарты для повторного открытия и редактирования.",
            "Модель должна быть расширяемой без регулярной переработки прикладного кода.",
        ],
    )
    doc.add_paragraph(
        "Исходя из указанных требований, в проекте выбрана реляционная модель данных PostgreSQL "
        "с локальной денормализацией на уровне таблицы снимков (JSONB)."
    )

    doc.add_heading("3.3 Концептуальная модель данных", level=1)
    doc.add_paragraph(
        "Концептуально данные разделены на пять взаимосвязанных доменных групп:"
    )
    bullets(
        doc,
        [
            "Нормативная группа: control_methods, regulatory_documents.",
            "Группа сварных соединений и схем: type_of_welded_joint, params_by_type_welding_joint, cheme_control, weld_type_to_scheme, width_of_controlled_area.",
            "Группа материалов и источников: type_metall, voltage_tube, class_film_radiation_source, rengen_apparatus, radiographic_film, control_sensitivity.",
            "Группа шаблонов операций: list_operation, operation_param.",
            "Группа пользовательских данных: tech_cards.",
        ],
    )
    doc.add_paragraph(
        "Такое разделение обеспечивает независимую эволюцию справочников, расчетных таблиц и "
        "пользовательских артефактов."
    )

    doc.add_heading("3.4 Логическая модель и ключевые сущности", level=1)
    doc.add_heading("3.4.1 Нормативная иерархия", level=2)
    bullets(
        doc,
        [
            "control_methods(id, name) — корневой справочник методов контроля.",
            "regulatory_documents(id, control_methods_id, name) — документы, привязанные к методу.",
            "type_of_welded_joint(id, regulatory_documents_id, name, image_ref) — типы швов в контексте выбранного документа.",
        ],
    )
    doc.add_paragraph(
        "Иерархия реализует последовательный фильтр вариантов на пользовательском интерфейсе "
        "и предотвращает выбор методологически несовместимых комбинаций."
    )

    doc.add_heading("3.4.2 Динамические параметры шва и схемы контроля", level=2)
    bullets(
        doc,
        [
            "params_by_type_welding_joint(id, type_of_welded_joint_id, name, subtitle) — описание параметров, добавляемых в карточку динамически.",
            "cheme_control(id, name, image_ref) — справочник схем просвечивания/контроля.",
            "weld_type_to_scheme(id, type_of_welded_joint_id, cheme_control_id) — ассоциативная таблица many-to-many.",
            "width_of_controlled_area(id, type_of_welded_joint_id, min_nominal_diametr, max_nominal_diam, val) — технологические правила по ширине зоны.",
        ],
    )
    doc.add_paragraph(
        "Критически важный момент модели — таблица weld_type_to_scheme. Она позволяет "
        "расширять варианты схем без изменения структуры прикладных алгоритмов: "
        "достаточно добавить новые строки связей."
    )

    doc.add_heading("3.4.3 Материалы, оборудование и параметры источника", level=2)
    bullets(
        doc,
        [
            "type_metall(id, material, id_standard) — справочник материалов.",
            "voltage_tube(id, type_metall_id, radiation_thickness, voltage_tube) — табличные ограничения по максимально допустимому напряжению.",
            "rengen_apparatus(id, name, voltage_on_tube, val, focal_spot_size) — список аппаратов и их технических характеристик.",
            "class_film_radiation_source(id, type_metall_id, min_radiation_thickness, max_radiation_thickness, class_a, class_b, class_c) — класс пленки по толщине и материалу.",
            "radiographic_film(id, film_class, name) — конкретные марки/виды пленки.",
            "control_sensitivity(id, thickness_from, thickness_to, class_a, class_b, class_c) — чувствительность контроля по диапазонам толщин.",
        ],
    )
    doc.add_paragraph(
        "Связка type_metall + voltage_tube + rengen_apparatus формирует один из самых "
        "важных алгоритмов системы: выбор допустимого ИИИ по нормативному пределу напряжения."
    )

    doc.add_heading("3.4.4 Операционные шаблоны и снимки карты", level=2)
    bullets(
        doc,
        [
            "list_operation(id, name) — именованные наборы операций.",
            "operation_param(id, id_list, val, val2, name_param) — элементы набора операций.",
            "tech_cards(id, name, card_data JSONB, created_at, updated_at) — сохраненные снимки пользовательских карт.",
        ],
    )
    doc.add_paragraph(
        "Использование JSONB в tech_cards является осознанным архитектурным решением: "
        "оно фиксирует полное состояние карточки в момент сохранения и допускает развитие "
        "структуры UI-полей без непрерывного изменения реляционной схемы."
    )

    doc.add_heading("3.5 Физическая реализация в PostgreSQL", level=1)
    bullets(
        doc,
        [
            "Для сущностей применены surrogate keys (serial / bigserial).",
            "Целостность обеспечивается внешними ключами между всеми ключевыми доменными таблицами.",
            "Используются типы text, integer, double precision, timestamp with time zone, jsonb.",
            "Временные поля в tech_cards поддерживают аудит изменений.",
            "Отказ от каскадных удалений (NO ACTION) повышает сохранность нормативных данных.",
        ],
    )
    doc.add_paragraph(
        "Выбор PostgreSQL обоснован поддержкой JSONB, надежной транзакционностью и "
        "возможностью дальнейшего масштабирования за счет индексов, представлений и процедур."
    )

    doc.add_heading("3.6 Анализ кардинальностей и референциальной целостности", level=1)
    nums(
        doc,
        [
            "control_methods (1) -> (N) regulatory_documents.",
            "regulatory_documents (1) -> (N) type_of_welded_joint.",
            "type_of_welded_joint (1) -> (N) params_by_type_welding_joint.",
            "type_of_welded_joint (N) <-> (N) cheme_control (через weld_type_to_scheme).",
            "type_metall (1) -> (N) voltage_tube.",
            "type_metall (1) -> (N) class_film_radiation_source.",
            "list_operation (1) -> (N) operation_param.",
        ],
    )
    doc.add_paragraph(
        "Указанные зависимости обеспечивают корректную навигацию по предметной модели и "
        "служат основой для каскадных выборок на уровне прикладных сервисов."
    )

    doc.add_heading("3.7 Интересные инженерные особенности модели", level=1)
    bullets(
        doc,
        [
            "Данные выступают как «табличный движок правил»: значительная часть логики не зашита в код, а хранится в БД.",
            "Гибкое связывание швов и схем через ассоциативную таблицу избавляет от жестких enum-ограничений в приложении.",
            "params_by_type_welding_joint реализует data-driven формирование UI, что снижает стоимость добавления новых типов швов.",
            "Разделение class_film_radiation_source и radiographic_film задает двухэтапный выбор: класс -> конкретная пленка.",
            "JSONB-хранилище tech_cards позволяет хранить историю реальных пользовательских сценариев без потери контекста.",
        ],
    )

    doc.add_heading("3.8 Риски и направления оптимизации", level=1)
    bullets(
        doc,
        [
            "Требуются дополнительные UNIQUE-ограничения в справочниках для предотвращения дублирования записей.",
            "Желательны индексы по полям lookup-запросов: type_metall(material), voltage_tube(type_metall_id, radiation_thickness).",
            "Для ускорения аналитики по tech_cards целесообразны GIN-индексы по card_data.",
            "Стоит добавить триггер автозаполнения updated_at при изменении записи tech_cards.",
            "Для нормативных таблиц целесообразно версионирование (valid_from/valid_to) и аудит изменений.",
        ],
    )

    doc.add_heading("3.9 Выводы по главе", level=1)
    doc.add_paragraph(
        "В результате проектирования сформирована целостная и расширяемая модель данных, "
        "которая покрывает как нормативный контур предметной области, так и эксплуатационный "
        "контур пользовательских данных. База данных реализует инженерно обоснованный баланс "
        "между строгой реляционной структурой и гибкостью хранения динамического состояния карт. "
        "Принятые решения позволяют масштабировать систему, развивать бизнес-логику без "
        "критических миграций и обеспечивать методическую корректность автоматических расчетов."
    )

    doc.save(OUT_PATH)
    print(OUT_PATH)


if __name__ == "__main__":
    main()
