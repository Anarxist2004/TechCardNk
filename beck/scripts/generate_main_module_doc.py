"""
Генерация Word-документа: модуль точки входа main.py.

Запуск из корня beck: python scripts/generate_main_module_doc.py
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "main_module_detailed.docx"


def _set_body_style(doc: Document) -> None:
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    pf = style.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_after = Pt(6)


def _p(doc: Document, text: str) -> None:
    doc.add_paragraph(text)


def _heading(doc: Document, text: str, level: int) -> None:
    doc.add_heading(text, level=level)


def build_document() -> Document:
    doc = Document()
    _set_body_style(doc)

    doc.add_heading("Модуль точки входа приложения (main.py)", 0)
    _p(
        doc,
        f"Документ составлен по файлу `main.py` в корне проекта `beck`. "
        f"Дата: {datetime.now().strftime('%d.%m.%Y %H:%M')}.",
    )

    _heading(doc, "1. Назначение модуля", 1)
    _p(
        doc,
        "Файл `main.py` выполняет роль **корня композиции** (composition root): здесь "
        "создаются конкретные реализации слоёв, собирается конвейер преобразований техкарты "
        "и запускается веб-адаптер. В нём нет предметной бизнес-логики техкарты — только "
        "импорты, константа подключения к БД, функция сборки `PipeLine` и функция `main`.",
    )

    _heading(doc, "2. Настройка вывода в консоль", 1)
    _p(
        doc,
        "После блока импортов выполняется присваивание "
        "`sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding=\"utf-8\")`. "
        "Это принудительно задаёт для стандартного вывода кодировку UTF-8, что уменьшает "
        "риск «кракозябр» в консоли Windows при `print` из репозитория и адаптера. "
        "Побочный эффект — обёртка вместо «сырого» stdout; для сложных сценариев "
        "(перенаправление, тесты) иногда требуют отдельной настройки.",
    )

    _heading(doc, "3. Константа DB_DSN", 1)
    _p(
        doc,
        "Строка `DB_DSN` задаёт параметры подключения `psycopg2` к PostgreSQL "
        "(хост, порт, имя базы, пользователь, пароль). В учебном и дипломном проекте "
        "такое хранение допустимо для простоты запуска; в реальной эксплуатации "
        "рекомендуется вынести секреты в переменные окружения или конфигурацию, "
        "не хранящуюся в системе контроля версий.",
    )

    _heading(doc, "4. Функция create_pipeline", 1)
    _p(
        doc,
        "Принимает экземпляр `PostgresDataBase` (репозиторий) и возвращает настроенный "
        "`PipeLine`. Все преобразователи добавляются вызовом `addChanger(..., 0)` — второй "
        "аргумент в данном проекте везде ноль (приоритет/фаза в очереди pipeline, "
        "согласно реализации `PipeLine`). Порядок добавления задаёт **порядок прохода** "
        "changer’ов при обновлении техкарты.",
    )
    _heading(doc, "4.1. Перечень changer’ов в порядке регистрации", 2)
    changers = [
        ("ControlMethodsFromDb", "Справочник методов контроля из БД."),
        ("RegulatoryDocumentsFromDb", "Нормативные документы по методу."),
        ("TypeOfWeldedJointFromDb", "Типы сварных соединений."),
        ("WeldedJointDiagramFromDb", "Схема/чертеж соединения (image_ref и связанные поля)."),
        ("ControlSchemesFromJointTypeDb", "Схемы контроля для выбранного типа соединения."),
        ("MaterialsFromDb", "Материалы из справочника."),
        ("RengenApparatusFromDb", "Рентген-аппараты; в конструктор передаётся `repos` дважды согласно сигнатуре класса."),
        ("ParamsByWeldedJointFromDb", "Доп. параметры по типу соединения."),
        ("CircButtWeldBeadFromThicknessTable", "Табличные данные по ширине зоны/условиям — без БД."),
        ("ControlSensitivityChanger", "Чувствительность контроля."),
        ("RengenApparatusForPanoramicScheme", "Параметры аппарата под панорамную схему."),
        ("RengenDistanceForDoubleWallScheme", "Расстояния для двухстеночной схемы."),
        ("RengenDistanceForEllipseFrontalScheme", "Расстояния для эллиптической фронтальной схемы."),
        ("RectilinearDrDistance", "Прямолинейная DR / расчёты расстояния."),
        ("SensitivityEtalonByMaterial", "Эталон чувствительности по материалу."),
        ("RadiographicFilmFromDb", "Радиографическая плёнка из БД."),
        ("IntensifyingScreenByVoltage", "Усиливающий экран по напряжению."),
        ("ProtectiveScreenByVoltage", "Защитный экран по напряжению."),
        ("ControlledZoneWidthStub", "Заглушка ширины контролируемой зоны."),
    ]
    for i, (name, desc) in enumerate(changers, start=1):
        _p(doc, f"{i}. {name} — {desc}")

    _heading(doc, "5. Функция main", 1)
    _p(
        doc,
        "Последовательность сборки: "
        "(1) `repos = PostgresDataBase(DB_DSN)` — подключение к БД; "
        "(2) `service = TechCardService(repos, create_pipeline(repos))` — сервис получает "
        "репозиторий и готовый pipeline; "
        "(3) `controller = ControllerWeb(service)` — тонкий контроллер над сервисом; "
        "(4) `create_adapter(controller)` — регистрация маршрутов FastAPI и **блокирующий** "
        "запуск `uvicorn` на `0.0.0.0:8000` (см. `adapterWeb.py`). После вызова `create_adapter` "
        "управление не возвращается в `main` до остановки сервера.",
    )

    _heading(doc, "6. Запуск как сценария", 1)
    _p(
        doc,
        "Блок `if __name__ == \"__main__\": main()` позволяет выполнять `python main.py` "
        "из корня проекта: при импорте `main` как модуля из другого кода функция `main` "
        "автоматически не вызывается.",
    )

    _heading(doc, "7. Зависимости модуля (обзор)", 1)
    _p(
        doc,
        "Импорты охватывают слой репозитория (`PostgresDataBase`), контроллер и адаптер "
        "(`ControllerWeb`, `create_adapter`), ядро сервиса (`PipeLine`, `TechCardService`) "
        "и полный набор changer-модулей из `services/Changers/`. Таким образом, `main.py` "
        "является единственным местом, где перечислен **полный** состав конвейера для "
        "данной конфигурации приложения.",
    )

    return doc


def main() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = build_document()
    doc.save(OUTPUT)
    print(f"Written: {OUTPUT}")


if __name__ == "__main__":
    main()
