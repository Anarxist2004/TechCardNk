"""
Генерация дневников производственной и преддипломной практики по шаблону
«Дневник учебной практики.docx» (та же вёрстка, новые даты и содержание).

Запуск из корня beck: python scripts/generate_production_prediploma_diaries.py
"""

from __future__ import annotations

from datetime import date, timedelta
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
TEMPLATE = DOCS / "Дневник учебной практики.docx"
OUT_PROD = DOCS / "Дневник производственной практики.docx"
OUT_PRE = DOCS / "Дневник преддипломной практики.docx"

# Число строк данных в таблице (без шапки) — как в шаблоне
DATA_ROWS = 18

# Интервалы практик (границы включительно; выходные исключаются)
PROD_START = date(2026, 3, 2)
PROD_END = date(2026, 4, 11)
PRE_START = date(2026, 4, 13)
PRE_END = date(2026, 5, 8)


def _weekdays_between(start: date, end: date) -> list[date]:
    out: list[date] = []
    d = start
    while d <= end:
        if d.weekday() < 5:
            out.append(d)
        d += timedelta(days=1)
    return out


def _pick_evenly(days: list[date], n: int) -> list[date]:
    """Равномерная выборка n дат из отсортированного списка рабочих дней."""
    if len(days) < n:
        raise ValueError(
            f"В интервале только {len(days)} рабочих дней, нужно не меньше {n}"
        )
    if len(days) == n:
        return list(days)
    indices = [round(i * (len(days) - 1) / (n - 1)) for i in range(n)]
    return [days[i] for i in indices]


def _fmt_d(d: date) -> str:
    return d.strftime("%d.%m.%y")


def _fill_from_template(
    out_path: Path,
    practice_genitive: str,
    entries: list[tuple[str, str]],
) -> None:
    if not TEMPLATE.exists():
        raise FileNotFoundError(f"Шаблон не найден: {TEMPLATE}")

    doc = Document(str(TEMPLATE))

    for p in doc.paragraphs:
        if "учебной" in p.text and "практики" in p.text:
            p.text = p.text.replace("учебной", practice_genitive)

    table = doc.tables[0]
    n = min(len(entries), DATA_ROWS)
    for i in range(n):
        row = table.rows[i + 1]
        date_s, desc = entries[i]
        row.cells[0].text = date_s
        row.cells[1].text = desc
        row.cells[2].text = ""
    for i in range(n, DATA_ROWS):
        row = table.rows[i + 1]
        row.cells[0].text = ""
        row.cells[1].text = ""
        row.cells[2].text = ""

    # Таблица подписей: ФИО студента
    sig = doc.tables[1]
    if len(sig.rows) >= 2:
        cell = sig.rows[1].cells[2]
        cell.text = cell.text.replace("Григораш Д.Ю.", "Марков Д.С.")

    doc.save(str(out_path))


def main() -> None:
    dates_prod = _pick_evenly(_weekdays_between(PROD_START, PROD_END), DATA_ROWS)
    dates_pre = _pick_evenly(_weekdays_between(PRE_START, PRE_END), DATA_ROWS)

    prod_descriptions = [
        "Вводный инструктаж по охране труда и информационной безопасности; ознакомление с подразделением",
        "Анализ предметной области: операционная технологическая карта радиографического контроля сварных соединений",
        "Обзор нормативных документов и требований к оформлению и содержанию ОТК РК",
        "Проектирование логической структуры базы данных PostgreSQL (справочники, техкарты, параметры операций)",
        "Реализация слоя репозитория: методы чтения справочников и связанных сущностей",
        "Реализация сохранения и загрузки снапшотов техкарты; сериализация полей в JSON",
        "Разработка сервисного слоя и конвейера преобразований (pipeline) структуры карты",
        "Реализация преобразователей данных (changers) для зависимых полей техкарты",
        "Интеграция расчётных правил: схемы просвечивания, расстояния, параметры контроля",
        "Разработка HTTP-слоя на FastAPI: маршруты, тело запросов, обработка ошибок",
        "Настройка взаимодействия клиента и сервера; проверка CORS и формата обмена JSON",
        "Отладка пользовательских сценариев заполнения и пересчёта техкарты",
        "Анализ производительности SQL-запросов; обработка сбоев подключения к СУБД",
        "Подготовка описания программных интерфейсов и структуры модулей для отчётности",
        "Комплексное тестирование с тестовой базой данных PostgreSQL",
        "Разбор промежуточной версии с руководителем практики от организации; фиксация замечаний",
        "Внесение исправлений и локальный рефакторинг модулей по результатам проверки",
        "Оформление отчёта по производственной практике; подготовка приложений к проверке",
    ]

    pre_descriptions = [
        "Согласование индивидуального плана преддипломной практики с научным руководителем",
        "Уточнение архитектуры серверной части приложения; актуализация диаграмм компонентов",
        "Доработка и проверка преобразователей техкарты (changers) по замечаниям руководителя",
        "Разработка модульных тестов (pytest) для модели данных, конвейера и сервисного слоя",
        "Проверка расчётных и табличных алгоритмов на контрольных примерах; фиксация результатов",
        "Интеграционные проверки с базой данных; сверка данных справочников и сохранённых техкарт",
        "Сопоставление реализованного функционала с требованиями технического задания",
        "Подготовка иллюстраций к пояснительной записке: последовательности, структура подсистем",
        "Оформление раздела пояснительной записки «Тестирование»; описание набора автоматических тестов",
        "Оформление разделов по трудоёмкости и этапам жизненного цикла разработки",
        "Подготовка текстовых отчётов по подсистемам (генерация документов Word по коду проекта)",
        "Репетиция защиты: сценарий демонстрации и презентация результатов",
        "Устранение дефектов, выявленных при предварительной приёмке на кафедре",
        "Консультация по оформлению графического и текстового материала выпускной квалификационной работы",
        "Финальная вычитка пояснительной записки; проверка ссылок и оформления по методическим указаниям",
        "Сбор сведений для отзыва и рецензии; согласование формулировок с руководителем",
        "Передача электронных материалов в выпускающую кафедру; оформление дневника практики",
        "Подписание дневника преддипломной практики; завершение отчётности по практике",
    ]

    assert len(prod_descriptions) == DATA_ROWS
    assert len(pre_descriptions) == DATA_ROWS

    prod_entries = list(zip((_fmt_d(d) for d in dates_prod), prod_descriptions))
    pre_entries = list(zip((_fmt_d(d) for d in dates_pre), pre_descriptions))

    _fill_from_template(OUT_PROD, "производственной", prod_entries)
    _fill_from_template(OUT_PRE, "преддипломной", pre_entries)

    print(f"Written: {OUT_PROD}")
    print(f"Written: {OUT_PRE}")


if __name__ == "__main__":
    main()
