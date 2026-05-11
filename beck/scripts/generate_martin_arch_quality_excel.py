from datetime import datetime
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference, ScatterChart, Series
from openpyxl.styles import Alignment, Font

CURRENT_DIR = Path(__file__).resolve().parent
if str(CURRENT_DIR) not in sys.path:
    sys.path.insert(0, str(CURRENT_DIR))

from generate_martin_arch_quality_doc import (
    TARGET_PACKAGES,
    collect_modules,
    compute_metrics,
    layer_violations,
    package_cycles,
    score,
)


OUT_PATH = r"e:\sem8\ful\Diplom_NK\beck\docs\architecture_quality_martin.xlsx"


def _autosize_columns(ws):
    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            val = "" if cell.value is None else str(cell.value)
            max_len = max(max_len, len(val))
        ws.column_dimensions[col_letter].width = min(42, max(10, max_len + 2))


def build_workbook() -> Workbook:
    modules = collect_modules()
    metrics = compute_metrics(modules)
    violations = layer_violations(modules)
    cycles = package_cycles(modules)
    scores = score(metrics, violations, cycles)

    wb = Workbook()
    ws = wb.active
    ws.title = "Metrics"

    ws["A1"] = "Расчет качества архитектуры по Роберту Мартину"
    ws["A2"] = f"Дата расчета: {datetime.now().strftime('%d.%m.%Y %H:%M')}"
    ws["A1"].font = Font(size=14, bold=True)
    ws["A2"].font = Font(italic=True)

    headers = ["Слой", "Nc", "Na", "Ca", "Ce", "A", "I", "D"]
    for idx, h in enumerate(headers, 1):
        c = ws.cell(row=4, column=idx, value=h)
        c.font = Font(bold=True)
        c.alignment = Alignment(horizontal="center")

    row = 5
    for pkg in TARGET_PACKAGES:
        m = metrics[pkg]
        ws.cell(row=row, column=1, value=pkg)
        ws.cell(row=row, column=2, value=m["Nc"])
        ws.cell(row=row, column=3, value=m["Na"])
        ws.cell(row=row, column=4, value=m["Ca"])
        ws.cell(row=row, column=5, value=m["Ce"])
        ws.cell(row=row, column=6, value=round(m["A"], 4))
        ws.cell(row=row, column=7, value=round(m["I"], 4))
        ws.cell(row=row, column=8, value=round(m["D"], 4))
        row += 1

    ws["A10"] = "Индекс качества (агрегированный)"
    ws["A10"].font = Font(bold=True)
    ws["A11"] = "Main Sequence Score"
    ws["B11"] = round(scores["main_seq_score"], 2)
    ws["A12"] = "Layering Score"
    ws["B12"] = round(scores["layer_score"], 2)
    ws["A13"] = "Cycle Score"
    ws["B13"] = round(scores["cycle_score"], 2)
    ws["A14"] = "Architecture Quality Index"
    ws["B14"] = round(scores["total"], 2)
    ws["A14"].font = Font(bold=True)
    ws["B14"].font = Font(bold=True)

    ws["D10"] = "Нарушения слоев"
    ws["D10"].font = Font(bold=True)
    if violations:
        for i, v in enumerate(violations, 11):
            ws.cell(row=i, column=4, value=v)
    else:
        ws["D11"] = "Не обнаружено"

    ws["F10"] = "Циклы слоев"
    ws["F10"].font = Font(bold=True)
    if cycles:
        for i, c in enumerate(cycles, 11):
            ws.cell(row=i, column=6, value=c)
    else:
        ws["F11"] = "Не обнаружено"

    # Подготовка данных для главного графика Martin (Main Sequence).
    ws["J4"] = "A_actual"
    ws["K4"] = "I_actual"
    ws["L4"] = "A_ideal"
    ws["M4"] = "I_ideal=1-A"
    ws["N4"] = "Layer"
    for idx, pkg in enumerate(TARGET_PACKAGES, start=5):
        ws.cell(row=idx, column=10, value=round(metrics[pkg]["A"], 4))
        ws.cell(row=idx, column=11, value=round(metrics[pkg]["I"], 4))
        ws.cell(row=idx, column=14, value=pkg)

    # Линия главной последовательности.
    for i, a in enumerate([x / 10 for x in range(0, 11)], start=5):
        ws.cell(row=i, column=12, value=a)
        ws.cell(row=i, column=13, value=round(1 - a, 4))

    # График 1: Main Sequence (Scatter).
    scatter = ScatterChart()
    scatter.title = "Main Sequence (A vs I)"
    scatter.x_axis.title = "A (Abstractness)"
    scatter.y_axis.title = "I (Instability)"
    scatter.width = 14
    scatter.height = 8

    x_actual = Reference(ws, min_col=10, min_row=5, max_row=4 + len(TARGET_PACKAGES))
    y_actual = Reference(ws, min_col=11, min_row=5, max_row=4 + len(TARGET_PACKAGES))
    s_actual = Series(y_actual, x_actual, title="Слои проекта")
    s_actual.marker.symbol = "circle"
    s_actual.marker.size = 8
    s_actual.graphicalProperties.line.noFill = True
    scatter.series.append(s_actual)

    x_ideal = Reference(ws, min_col=12, min_row=5, max_row=15)
    y_ideal = Reference(ws, min_col=13, min_row=5, max_row=15)
    s_ideal = Series(y_ideal, x_ideal, title="Идеальная линия I=1-A")
    s_ideal.marker.symbol = "none"
    scatter.series.append(s_ideal)

    ws.add_chart(scatter, "A17")

    # График 2: Distance D по пакетам.
    bar_d = BarChart()
    bar_d.type = "col"
    bar_d.title = "Distance from Main Sequence (D)"
    bar_d.y_axis.title = "D"
    bar_d.x_axis.title = "Слой"
    data_d = Reference(ws, min_col=8, min_row=4, max_row=4 + len(TARGET_PACKAGES))
    cats = Reference(ws, min_col=1, min_row=5, max_row=4 + len(TARGET_PACKAGES))
    bar_d.add_data(data_d, titles_from_data=True)
    bar_d.set_categories(cats)
    bar_d.width = 10
    bar_d.height = 6
    ws.add_chart(bar_d, "I17")

    # График 3: Coupling (Ca vs Ce).
    bar_c = BarChart()
    bar_c.type = "col"
    bar_c.grouping = "clustered"
    bar_c.overlap = 0
    bar_c.title = "Layer Coupling (Ca vs Ce)"
    bar_c.y_axis.title = "Количество зависимостей"
    data_c = Reference(ws, min_col=4, max_col=5, min_row=4, max_row=4 + len(TARGET_PACKAGES))
    bar_c.add_data(data_c, titles_from_data=True)
    bar_c.set_categories(cats)
    bar_c.width = 10
    bar_c.height = 6
    ws.add_chart(bar_c, "I30")

    # График 4: Сводный индекс.
    ws["A30"] = "Показатель"
    ws["B30"] = "Баллы"
    ws["A31"] = "Main Sequence Score"
    ws["B31"] = round(scores["main_seq_score"], 2)
    ws["A32"] = "Layering Score"
    ws["B32"] = round(scores["layer_score"], 2)
    ws["A33"] = "Cycle Score"
    ws["B33"] = round(scores["cycle_score"], 2)
    ws["A34"] = "Architecture Quality Index"
    ws["B34"] = round(scores["total"], 2)

    line = LineChart()
    line.title = "Сводный профиль качества"
    line.y_axis.title = "Баллы"
    d = Reference(ws, min_col=2, min_row=30, max_row=34)
    c = Reference(ws, min_col=1, min_row=31, max_row=34)
    line.add_data(d, titles_from_data=True)
    line.set_categories(c)
    line.width = 14
    line.height = 6
    ws.add_chart(line, "A36")

    _autosize_columns(ws)
    return wb


def main() -> None:
    wb = build_workbook()
    wb.save(OUT_PATH)
    print(OUT_PATH)


if __name__ == "__main__":
    main()
