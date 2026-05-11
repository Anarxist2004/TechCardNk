"""
Генерация Word-документа: слой репозитория (интерфейсы и реализация PostgresDataBase).

Запуск из корня beck: python scripts/generate_repository_layer_doc.py
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "repository_layer_detailed.docx"


def _set_body_style(doc: Document) -> None:
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    pf = style.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_after = Pt(6)


def _p(doc: Document, text: str, *, bullet: bool = False) -> None:
    doc.add_paragraph(text, style="List Bullet" if bullet else None)


def _heading(doc: Document, text: str, level: int) -> None:
    doc.add_heading(text, level=level)


def build_document() -> Document:
    doc = Document()
    _set_body_style(doc)

    doc.add_heading(
        "Слой репозитория: интерфейсы доступа к данным и реализация на PostgreSQL",
        0,
    )
    _p(
        doc,
        (
            "Документ сформирован автоматически по исходному коду каталога "
            "`repositories/`. "
            f"Дата: {datetime.now().strftime('%d.%m.%Y %H:%M')}. "
            "Описаны назначение слоя, все абстрактные контракты в "
            "`repositories/Interfaces/` и то, как они реализованы в классе "
            "`PostgresDataBase` (`repositories/postgresDataBase.py`)."
        ),
    )

    _heading(doc, "1. Назначение и устройство слоя", 1)
    _p(
        doc,
        "Слой репозитория отделяет бизнес-логику и HTTP-слой от деталей хранения данных. "
        "Сервисы и преобразователи (changers) зависят от абстрактных интерфейсов "
        "(`IRepository`, `IControlMethodsDB` и др.), а не от конкретного драйвера БД. "
        "В проекте используется единственная реализация — класс `PostgresDataBase`, "
        "который наследует все перечисленные интерфейсы и выполняет SQL к схеме `public` "
        "через библиотеку `psycopg2` с курсором `RealDictCursor` (строки результата — словари).",
    )
    _p(
        doc,
        "При успешном подключении в конструкторе выводится сообщение об успехе; при ошибке "
        "подключения `conn` и `cursor` становятся `None`. Большинство методов чтения в этом "
        "случае возвращают пустой список, `None` или `0.0` и не бросают исключение. "
        "Исключение `RuntimeError` выбрасывается при попытке сохранить снапшот техкарты "
        "без инициализированного курсора. Метод `save_tech_card_snapshot` при ошибке SQL "
        "выполняет `rollback` и пробрасывает исключение дальше.",
    )
    _p(
        doc,
        "Дублирование контракта: метод `get_operation_params_by_list_id` объявлен и в "
        "`IRepository`, и в `IOperationParamDB`; в `PostgresDataBase` реализована одна "
        "общая функция, удовлетворяющая обоим интерфейсам.",
    )

    _heading(doc, "2. Класс PostgresDataBase: состав интерфейсов", 1)
    _p(
        doc,
        "`PostgresDataBase` объявлен как наследник: "
        "`IRepository[TechCardData]`, `IControlMethodsDB`, `IRegulatoryDocumentsDB`, "
        "`ITypeOfWeldedJointDB`, `IParamsByTypeWeldingJointDB`, `IChemeControlDB`, "
        "`IMaterialsDB`, `IMaterialStandardDB`, `IRengenApparatusDB`, "
        "`IRadiographicFilmDB`, `IOperationParamDB`, `IVoltageTubeDB`. "
        "Других классов, реализующих эти интерфейсы в кодовой базе, не обнаружено.",
    )

    _heading(doc, "3. IRepository[T] — файл i_repository.py", 1)
    _p(
        doc,
        "Обобщённый интерфейс «репозитория техкарты» с тип-параметром `T`. В проекте "
        "`PostgresDataBase` параметризован как `IRepository[TechCardData]`, однако методы "
        "сохранения и загрузки работают с `Dict[str, Any]` для сериализуемого снапшота "
        "и метаданных записи в таблице `public.tech_cards`.",
    )
    _heading(doc, "3.1. Абстрактные методы", 2)
    for t in (
        "`get_operation_params_by_list_id(list_id)` — параметры операций для блока "
        "«перечень операций РК» (см. также `IOperationParamDB`).",
        "`save_tech_card_snapshot(name, card_data, card_id=None)` — вставка или обновление "
        "снапшота: при переданном `card_id` выполняется `UPDATE ... RETURNING`; если строка "
        "не затронута, выполняется `INSERT`. Поле `card_data` передаётся как `psycopg2.extras.Json`. "
        "После успеха — `commit`. Даты `created_at`/`updated_at` в ответе приводятся к ISO-строке "
        "через вспомогательный `_serialize_dt`.",
        "`list_saved_tech_cards()` — список записей из `tech_cards` без поля `card_data`, "
        "сортировка по убыванию `updated_at`, затем `id`.",
        "`get_saved_tech_card(card_id)` — одна запись с полем `card_data` и сериализованными датами.",
    ):
        _p(doc, t, bullet=True)

    _heading(doc, "3.2. Реализация в PostgresDataBase", 2)
    _p(
        doc,
        "Конструктор вызывает `IRepository.__init__(self)`. Запрос параметров операций — "
        "соединение `operation_param` с `list_operation` по `id_list`, проекция полей "
        "`id`, `id_list`, `name_param`, `val`, `val2`, порядок по `op.id`.",
    )

    _heading(doc, "4. IControlMethodsDB — i_control_methods_db.py", 1)
    _p(doc, "Доступ к справочнику методов контроля (`public.control_methods`).")
    _p(
        doc,
        "`get_control_methods()` — выборка `id`, `name`, сортировка по `id`.",
        bullet=True,
    )

    _heading(doc, "5. IRegulatoryDocumentsDB — i_regulatory_documents_db.py", 1)
    _p(doc, "Доступ к нормативным документам (`public.regulatory_documents`).")
    for t in (
        "`get_regulatory_documents()` — все строки: `id`, `control_methods_id`, `name`.",
        "`get_regulatory_documents_by_control_method_id(control_method_id)` — фильтр по "
        "`control_methods_id`.",
        "`get_regulatory_documents_for_method_name(method_name)` — соединение с "
        "`control_methods`, сравнение имён без учёта регистра и краевых пробелов (`lower(btrim(...))`).",
    ):
        _p(doc, t, bullet=True)

    _heading(doc, "6. ITypeOfWeldedJointDB — i_type_of_welded_joint_db.py", 1)
    _p(doc, "Типы сварных соединений (`public.type_of_welded_joint`).")
    for t in (
        "`get_type_of_welded_joints()` — поля `id`, `regulatory_documents_id`, `name`, `image_ref`.",
        "`get_type_of_welded_joints_by_regulatory_document_id` — фильтр по внешнему ключу.",
        "`get_type_of_welded_joints_for_regulatory_document_name` — если строка из цифр, "
        "интерпретируется как id документа и делегирует методу по id; иначе join с "
        "`regulatory_documents` по совпадению имени (без учёта регистра).",
        "`get_welded_joint_image_ref(joint_name_or_id)` — по целочисленному id, по строке-числу "
        "или по имени (`lower(btrim(name))`), возврат `image_ref` или `None` при пустом значении.",
    ):
        _p(doc, t, bullet=True)

    _heading(doc, "7. IParamsByTypeWeldingJointDB — i_params_by_type_welding_joint_db.py", 1)
    _p(
        doc,
        "`get_params_by_welded_joint_type(joint_name_or_id)` — из таблицы "
        "`params_by_type_welding_joint` через join с `type_of_welded_joint`: условие либо "
        "`twj.id = %s`, либо равенство нормализованных имён. Возвращаются `id`, `name`, `subtitle` "
        "параметров, порядок по `p.id`.",
    )

    _heading(doc, "8. IChemeControlDB — i_cheme_control_db.py", 1)
    _p(
        doc,
        "`get_control_schemes_for_welded_joint(joint_name_or_id)` — схемы из "
        "`cheme_control`, связанные с типом соединения через `weld_type_to_scheme` "
        "(поля результата: `id`, `name`, `image_ref` схемы). Условие отбора типа соединения "
        "аналогично предыдущему интерфейсу (id или имя).",
    )

    _heading(doc, "9. IMaterialsDB — i_materials_db.py", 1)
    _p(
        doc,
        "`get_materials()` — все записи `type_metall`: `id`, `material`, сортировка по `id`.",
    )

    _heading(doc, "10. IMaterialStandardDB — i_material_standard_db.py", 1)
    _p(
        doc,
        "`get_material_standard_id(material_name_or_id)` — поле `id_standard` из "
        "`type_metall` по числовому id, строке-числу или по имени материала "
        "(нормализация `lower(btrim(material))`). Возврат `int` или `None`.",
    )

    _heading(doc, "11. IRengenApparatusDB — i_rengen_apparatus_db.py", 1)
    for t in (
        "`get_rengen_apparatus()` — все записи `rengen_apparatus`: "
        "`id`, `name`, `val`, `focal_spot_size`, `voltage_on_tube`.",
        "`get_rengen_apparatus_by_name_or_id` — по id или строке-числу; иначе по совпадению "
        "`name` или `val` (оба сравниваются с ключом, `coalesce(val,'')` для null-safe), "
        "первая запись по `id`.",
    ):
        _p(doc, t, bullet=True)

    _heading(doc, "12. IRadiographicFilmDB — i_radiographic_film_db.py", 1)
    for t in (
        "`get_radiographic_films()` — `radiographic_film`: `id`, `film_class`, `name`.",
        "`get_radiographic_films_by_class_range(min_class, max_class)` — фильтр "
        "`film_class` в диапазоне включительно, сортировка по классу и `id`.",
    ):
        _p(doc, t, bullet=True)

    _heading(doc, "13. IOperationParamDB — i_operation_param_db.py", 1)
    _p(
        doc,
        "`get_operation_params_by_list_id` — см. раздел 3; реализация общая с `IRepository`.",
    )

    _heading(doc, "14. IVoltageTubeDB — i_voltage_tube_db.py", 1)
    _p(
        doc,
        "`get_max_voltage_kv_for_material_and_thickness(material_label, radiation_thickness_mm)` — "
        "по нормализованному имени материала находится `type_metall.id`. При пустой метке, "
        "NaN, отрицательной толщине или отсутствии материала возвращается `0.0`. "
        "Далее из `voltage_tube` выбирается первая строка с `radiation_thickness >= t` "
        "(порядок по возрастанию толщины); если такой нет — берётся строка с максимальной "
        "`radiation_thickness` для данного `type_metall_id`. Значение поля `voltage_tube` "
        "приводится к `float`; при отсутствии данных — `0.0`.",
    )

    _heading(doc, "15. Вспомогательные детали реализации", 1)
    for t in (
        "Сообщения об ошибках SQL выводятся в консоль через `print` с префиксом имени метода.",
        "Пустые или пробельные строковые ключи для поиска обычно приводят к пустому результату "
        "без запроса к БД (где это явно проверено в коде).",
        "Для имён и текстовых полей широко используется сравнение `lower(btrim(...))` для "
        "устойчивости к регистру и пробелам.",
    ):
        _p(doc, t, bullet=True)

    return doc


def main() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = build_document()
    doc.save(OUTPUT)
    print(f"Written: {OUTPUT}")


if __name__ == "__main__":
    main()
